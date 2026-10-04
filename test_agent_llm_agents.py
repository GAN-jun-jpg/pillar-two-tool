# -*- coding: utf-8 -*-
"""其余 Agent 的云端能力测试。"""
import io
import json

import openpyxl

from Agent.agents import ChartAgent, DataAgent, PlannerAgent, ReviewAgent, TaxAgent
from Agent.llm import LLMBrain, MockGateway
from Agent.schemas import WorkflowState


class Upload:
    def __init__(self, data: bytes, name: str):
        self._data = data
        self.name = name

    def read(self) -> bytes:
        return self._data


def _workbook_bytes() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "利润表"
    ws.append(["科目", "2025", "单位：万元"])
    ws.append(["营业总收入", 50000, ""])
    ws.append(["利润总额", 10000, ""])
    ws.append(["所得税费用", 1500, ""])
    ws.append(["递延所得税费用", 200, ""])

    ws2 = wb.create_sheet("资产负债表")
    ws2.append(["科目", "2025", "单位：万元"])
    ws2.append(["固定资产", 20000, ""])
    ws2.append(["资产总计", 50000, ""])

    ws3 = wb.create_sheet("现金流量表")
    ws3.append(["科目", "2025", "单位：万元"])
    ws3.append(["经营活动产生的现金流量净额", 8000, ""])
    ws3.append(["投资活动产生的现金流量净额", -2000, ""])
    ws3.append(["支付给职工以及为职工支付的现金", 1200, ""])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _row(**kw):
    base = {
        "name": "测试辖区", "profit": 1000.0, "current_tax": 100.0,
        "deferred_tax": 0.0, "revenue": 5000.0, "payroll": 0.0,
        "tangible_assets": 0.0, "parent_idx": None, "ownership": 1.0,
        "qdmtt_applies": False, "utpr_applies": True, "dtl_ledger": [],
    }
    base.update(kw)
    return base


def test_planner_agent_llm_plan():
    gateway = MockGateway([
        '{"input_mode": "rows", "steps": ["validate", "calculate"]}'
    ])
    brain = LLMBrain(gateway=gateway, provider="deepseek")
    state = WorkflowState.new()
    PlannerAgent(brain=brain).run(
        state, rows=[_row()], include_charts=False, export_gir=False)
    assert state.metadata["plan_source"] == "llm"
    assert state.metadata["workflow_plan"] == ["validate", "calculate"]


def test_data_agent_llm_suggestions():
    gateway = MockGateway([
        '{"suggestions": [{"field": "payroll", "source": "cash_to_employees", "reason": "职工现金支出"}], "warnings": []}'
    ])
    brain = LLMBrain(gateway=gateway, provider="deepseek")
    state = WorkflowState.new()
    DataAgent(brain=brain).run(
        state,
        uploaded_file=Upload(_workbook_bytes(), "demo.xlsx"),
        jurisdiction_name="测试辖区",
    )
    assert state.mapped_rows[0]["payroll"] == 1200.0
    assert state.metadata["llm_data_suggestions"]["suggestions"][0]["field"] == "payroll"


def test_review_agent_llm_review():
    gateway = MockGateway([
        '{"decision": "pass", "summary": "校验通过", "issues": []}'
    ])
    brain = LLMBrain(gateway=gateway, provider="deepseek")
    state = WorkflowState.new()
    ReviewAgent(brain=brain).run(state, rows=[_row()], calc_year=2024)
    assert state.status != "failed"
    assert state.metadata["llm_review"]["decision"] == "pass"


def test_chart_agent_llm_plan():
    state = WorkflowState.new()
    rows = [_row()]
    TaxAgent().run(state, rows=rows, calc_year=2024)

    gateway = MockGateway([
        '{"chart_plan": ["ETR 柱状图"], "captions": {"etr": "ETR 对比"}, "report_outline": ["结论"]}'
    ])
    brain = LLMBrain(gateway=gateway, provider="deepseek")
    ChartAgent(include_charts=True, export_gir=False, brain=brain).run(
        state, rows=rows, group_name="测试集团")
    assert state.chart_data is not None
    assert state.metadata["llm_chart_plan"]["chart_plan"] == ["ETR 柱状图"]

def test_chart_agent_executes_llm_chart_plan():
    state = WorkflowState.new()
    rows = [_row()]
    TaxAgent().run(state, rows=rows, calc_year=2024)

    gateway = MockGateway([
        '{"chart_plan": [{"id": "etr_custom", "chart_type": "bar", "title": "自定义 ETR", "x": "name", "y": ["etr"]}], "captions": {"etr_custom": "ETR 对比"}, "report_outline": []}'
    ])
    brain = LLMBrain(gateway=gateway, provider="deepseek")
    ChartAgent(include_charts=True, export_gir=False, brain=brain).run(
        state, rows=rows, group_name="测试集团")
    assert "etr_custom" in state.chart_data
    assert state.metadata["chart_source"] == "llm_plan"
    assert state.chart_data["etr_custom"].layout.title.text == "自定义 ETR"

def test_chart_agent_accepts_aliases_and_alternate_keys():
    state = WorkflowState.new()
    rows = [_row()]
    TaxAgent().run(state, rows=rows, calc_year=2024)

    gateway = MockGateway([
        '{"charts": [{"id": "etr_alias", "chart_type": "bar_chart", "title": "ETR", "x": "jurisdiction", "y": ["effective_tax_rate"]}], "captions": {}, "report_outline": []}'
    ])
    brain = LLMBrain(gateway=gateway, provider="deepseek")
    ChartAgent(include_charts=True, export_gir=False, brain=brain).run(
        state, rows=rows, group_name="测试集团")
    assert state.metadata["chart_source"] == "llm_plan"
    assert "etr_alias" in state.chart_data

def test_chart_agent_handles_nested_fields_and_tax_flow_totals():
    state = WorkflowState.new()
    rows = [_row()]
    TaxAgent().run(state, rows=rows, calc_year=2024)
    plan = {
        "chart_plan": [
            {"chart_id": "chart_1", "chart_type": "bar", "title": "补税金额",
             "data_source": "tax_summary", "fields": {"x": "辖区名称", "y": "补税金额（万元）"}},
            {"chart_id": "chart_2", "chart_type": "pie", "title": "资金流向",
             "data_source": "tax_flow_totals", "fields": {"category": "资金流向", "value": "金额（万元）"}},
            {"chart_id": "chart_3", "chart_type": "stacked_bar", "title": "留存导出",
             "data_source": "tax_flow_totals", "fields": {"x": "辖区名称", "y": "金额（万元）"}},
            {"chart_id": "chart_4", "chart_type": "waterfall", "title": "分配瀑布",
             "data_source": "tax_flow_totals", "fields": {"category": "分配环节", "value": "金额（万元）"}},
        ],
        "captions": {},
        "report_outline": [],
    }
    gateway = MockGateway([json.dumps(plan, ensure_ascii=False)])
    brain = LLMBrain(gateway=gateway, provider="deepseek")
    ChartAgent(include_charts=True, export_gir=False, brain=brain).run(
        state, rows=rows, group_name="测试集团")
    assert state.metadata["chart_source"] == "llm_plan"
    # 方案 A 后的契约：chart_data = 本地标准图（含三类分析图）+ 云端规划的补充图
    keys = set(state.chart_data.keys())
    assert {"chart_1", "chart_2", "chart_3", "chart_4"} <= keys, "云端规划的图必须在"
    assert {"etr_bar", "waterfall", "sankey", "risk_matrix", "topup_bridge"} <= keys, \
        "本地标准图也必须由 ChartAgent 产出（方案 A）"
    assert "tax_flow_sankey" not in keys, \
        "已有标准 sankey 时不再重复生成第二张流向图"

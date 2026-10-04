# -*- coding: utf-8 -*-
"""风险因果链与分析报告测试。

口径：
- 因果链的数字必须与 calculator 的结果逐项一致；
- 报告必须按 P1–P8 固定模板输出：P3 覆盖全部辖区、P5 带上 QDMTT/IIR/UTPR 分配、
  P7 汇总算前校验与算后复核，未实现的能力标注「未启用」；
- Word 版必须把 Markdown 表格转成真正的 Word 表格，并按页插入分页符。
"""
import datetime
import io
import re
import zipfile

import pytest

from analysis_report import (
    build_report_docx, build_report_markdown, build_risk_chains, chain_markdown,
    markdown_to_docx,
)
from calculator import assess_jurisdiction, compute_tax_flow, run_allocation

ROWS = [
    {"name": "开曼群岛", "profit": 2000.0, "current_tax": 20.0, "deferred_tax": 0.0,
     "revenue": 0.0, "payroll": 0.0, "tangible_assets": 0.0, "parent_idx": None,
     "ownership": 1.0, "qdmtt_applies": False, "utpr_applies": True, "dtl_ledger": []},
    {"name": "德国", "profit": 800.0, "current_tax": 200.0, "deferred_tax": 30.0,
     "revenue": 800.0, "payroll": 1500.0, "tangible_assets": 3500.0,
     "parent_idx": None, "ownership": 1.0, "qdmtt_applies": False,
     "utpr_applies": True, "dtl_ledger": []},
    {"name": "亏损辖区", "profit": -500.0, "current_tax": 0.0, "deferred_tax": 0.0,
     "revenue": 0.0, "payroll": 0.0, "tangible_assets": 0.0, "parent_idx": None,
     "ownership": 1.0, "qdmtt_applies": False, "utpr_applies": True, "dtl_ledger": []},
]

# P1–P8 固定模板（顺序即页面顺序）
PAGES = [
    "## P1 执行摘要",
    "## P2 集团及数据概况",
    "## P3 辖区 ETR & Top-up Tax",
    "## P4 补税形成机制",
    "## P5 QDMTT · IIR · UTPR 补税分配",
    "## P6 重点税务风险",
    "## P7 数据与计算质量审查",
    "## P8 管理建议",
]


def _assess(rows):
    return [
        assess_jurisdiction(
            r["profit"], r["current_tax"], payroll=r.get("payroll", 0.0),
            tangible_assets=r.get("tangible_assets", 0.0),
            deferred_tax=r.get("deferred_tax", 0.0), dtl_ledger=r.get("dtl_ledger", []),
            calc_year=2024, revenue=r.get("revenue", 0.0),
        )
        for r in rows
    ]


def _results():
    return _assess(ROWS)


@pytest.fixture()
def results():
    return _results()


@pytest.fixture()
def allocation(results):
    return run_allocation(results, {i: None for i in range(len(ROWS))},
                          {i: True for i in range(len(ROWS))})


@pytest.fixture()
def tax_flow(results, allocation):
    return compute_tax_flow(results, allocation, ROWS)


@pytest.fixture()
def flow(results, tax_flow):
    return results, tax_flow


# ── 因果链 ──

def test_risk_chain_numbers_match_calculator(flow):
    results, tax_flow = flow
    chains = build_risk_chains(results, ROWS, tax_flow)
    cayman = next(c for c in chains if c["name"] == "开曼群岛")
    result = results[0]

    values = dict(cayman["steps"])
    assert values["GloBE 利润"] == f"{result['profit']:,.2f} 万元"
    assert values["Covered Taxes"] == f"{result['covered_taxes']:,.2f} 万元"
    assert values["ETR"] == f"{result['etr']:.1%}"
    assert values["超额利润"] == f"{result['adjusted_profit']:,.2f} 万元"
    assert values["补税额"] == f"{result['topup_tax']:,.2f} 万元"
    assert cayman["topup_tax"] == pytest.approx(result["topup_tax"])
    assert cayman["etr"] == pytest.approx(result["etr"])
    assert cayman["idx"] == 0, "报告 P5/P6 需要靠 idx 关联税源流向"


def test_risk_chain_orders_topup_first(flow):
    results, tax_flow = flow
    chains = build_risk_chains(results, ROWS, tax_flow)
    assert chains[0]["name"] == "开曼群岛", "需补税的辖区应排在最前"
    assert chains[0]["topup_tax"] > 0


def test_safe_harbour_chain_explains_exemption(flow):
    results, tax_flow = flow
    chains = build_risk_chains(results, ROWS, tax_flow)
    germany = next(c for c in chains if c["name"] == "德国")
    assert "安全港" in dict(germany["steps"])
    assert "豁免" in germany["conclusion"]
    assert germany["topup_tax"] == 0.0


def test_loss_jurisdiction_chain_marks_not_applicable(flow):
    results, tax_flow = flow
    chains = build_risk_chains(results, ROWS, tax_flow)
    loss = next(c for c in chains if c["name"] == "亏损辖区")
    assert "n/a" in dict(loss["steps"])["ETR"]
    assert "不进入补税计算" in loss["conclusion"]


def test_chain_markdown_uses_trigger_label_only_for_topup(flow):
    results, tax_flow = flow
    chains = build_risk_chains(results, ROWS, tax_flow)
    cayman = next(c for c in chains if c["name"] == "开曼群岛")
    germany = next(c for c in chains if c["name"] == "德国")
    assert "触发依据" in chain_markdown(cayman)
    assert "触发依据" not in chain_markdown(germany)


def test_risk_chains_handle_empty_results():
    assert build_risk_chains([], []) == []


# ── 报告：payload ──

def _payload(flow, allocation):
    results, tax_flow = flow
    return {
        "results": results,
        "rows": ROWS,
        "allocation": allocation,
        "tax_flow": tax_flow,
        "calc_year": 2024,
        "scenario_name": "演示方案",
        "result_source": "agent",
        "data_source": "上传的 Excel",
        "rule_version": "规则库修订号 3",
        "llm_tax_analysis": {
            "analysis": "开曼群岛有效税率偏低。",
            "highlights": ["开曼 ETR 1.0%"],
            "actions": ["评估开曼 QDMTT 适用性"],
        },
        "llm_result_review": {
            "assessment": "consistent",
            "summary": "与本地复核结论一致",
            "concerns": ["DTL 台账建议逐年核对"],
            "followups": ["补全实体架构"],
        },
        "llm_chart_plan": {"report_outline": ["集团概览", "风险辖区", "优化建议"]},
        "result_review": {
            "decision": "pass",
            "summary": "结果审查通过：5/5 项检查通过，0 个错误，0 个警告。",
            "checks": [{"check": "ETR 一致性", "status": "passed", "detail": "全部匹配"},
                       {"check": "税源流向闭合", "status": "passed", "detail": "已闭合"}],
            "errors": [],
            "warnings": [],
            "metrics": {"jurisdictions": len(results), "topup_jurisdictions": 1,
                        "total_topup_tax": 280.0, "trace_count": 12, "articles": 4},
        },
        "validation_report": {
            "findings": [{"severity": "warning", "code": "W015", "jurisdiction": "亏损辖区",
                          "field": "profit", "message": "GloBE 利润为负数"},
                         {"severity": "info", "code": "I001", "jurisdiction": "德国",
                          "field": "payroll", "message": "薪酬为 0"}],
        },
        "calc_maintenance": {"passed": True, "tested_rules_dir": "候选规则目录",
                             "rule_changes": []},
        "rule_impact": {"totals": {"topup_before": 250.0, "topup_after": 280.0,
                                   "topup_delta": 30.0}},
    }


# ── 报告：模板结构 ──

def test_report_follows_eight_page_template_in_order(flow, allocation):
    report = build_report_markdown(_payload(flow, allocation))
    assert report.startswith("# Pillar Two 全球最低税负分析报告")
    positions = [report.index(page) for page in PAGES]
    assert positions == sorted(positions), "P1–P8 必须按顺序出现"
    assert "## 附录：口径与免责声明" in report


def test_report_header_carries_run_metadata(flow, allocation):
    report = build_report_markdown(_payload(flow, allocation))
    for text in ["演示方案", "2024", "云端 Agent 流程", "规则库修订号 3",
                 "上传的 Excel", "云端仅提供解读"]:
        assert text in report, text


def test_report_uses_the_previously_dropped_report_outline(flow, allocation):
    """report_outline 此前没有任何出口，这里必须出现在报告里。"""
    report = build_report_markdown(_payload(flow, allocation))
    assert "1. 集团概览" in report
    assert "3. 优化建议" in report
    assert "评估开曼 QDMTT 适用性" in report      # 云端建议动作
    assert "本次运行未记录到云端调用失败" not in report  # 报告不写这句界面文案


def test_p3_lists_every_jurisdiction_with_judgement(flow, allocation):
    report = build_report_markdown(_payload(flow, allocation))
    p3 = report.split("## P3 辖区 ETR & Top-up Tax")[1].split("## P4")[0]
    assert "开曼群岛" in p3 and "德国" in p3 and "亏损辖区" in p3
    assert "需补税" in p3 and "安全港豁免" in p3 and "不适用（利润 ≤ 0）" in p3
    assert "280.00" in p3, "补税额应与计算结果一致"


def test_p4_renders_chain_only_for_topup_jurisdictions(flow, allocation):
    report = build_report_markdown(_payload(flow, allocation))
    p4 = report.split("## P4 补税形成机制")[1].split("## P5")[0]
    assert "开曼群岛" in p4 and "超额利润" in p4
    assert "德国" not in p4, "无须补税的辖区不进入补税形成机制页"
    assert "其余 2 个辖区无须补税" in p4


def test_p1_summarises_kpis_and_top_jurisdictions(flow, allocation):
    report = build_report_markdown(_payload(flow, allocation))
    p1 = report.split("## P1 执行摘要")[1].split("## P2")[0]
    assert "需补税合计 **280.00 万元**" in p1
    assert "重点辖区" in p1 and "开曼群岛" in p1
    assert "算前校验 🔴 0 阻断 · 🟡 1 警告 · 🔵 1 提示" in p1
    assert "算后结果复核 pass" in p1


def test_p2_shows_group_totals_and_architecture(flow, allocation):
    report = build_report_markdown(_payload(flow, allocation))
    p2 = report.split("## P2 集团及数据概况")[1].split("## P3")[0]
    assert "GloBE 利润合计" in p2 and "2,300.00" in p2   # 2000 + 800 - 500
    assert "最终母公司（UPE）" in p2
    assert "未做 GloBE 亏损跨辖区抵扣" in p2


def test_p7_counts_findings_and_result_review_checks(flow, allocation):
    report = build_report_markdown(_payload(flow, allocation))
    p7 = report.split("## P7 数据与计算质量审查")[1].split("## P8")[0]
    assert "🔴 0 阻断 · 🟡 1 警告 · 🔵 1 提示" in p7
    assert "W015" in p7 and "GloBE 利润为负数" in p7
    assert "结果审查通过：5/5 项检查通过" in p7
    assert "ETR 一致性" in p7 and "税源流向闭合" in p7
    assert "法规追溯 12 条，覆盖 4 个条款" in p7
    assert "候选规则回归测试" in p7 and "通过" in p7
    assert "250.00 → 280.00 万元" in p7


def test_p7_reads_rule_gap_blocking_from_gaps_key(flow, allocation):
    """规则缺口工具返回的是 gaps / informational_gaps，报告必须按这两个键计数。"""
    payload = _payload(flow, allocation)
    payload["rule_gap"] = {
        "gaps": [{"kind": "column_conflict", "target": "无形资产",
                  "globe_field": "tangible_assets", "status": "conflict",
                  "detail": "列名语义与命中的字段不符", "suggestion": "人工确认该列"}],
        "informational_gaps": [{"kind": "optional_field_unmatched", "target": "递延所得税",
                                "globe_field": "current_tax", "status": "missing",
                                "detail": "可选字段未匹配", "suggestion": "补充规则"}],
        "column_conflict_count": 1,
        "unrecognized_column_count": 0,
        "coverage": {"checked": 12, "unmatched": 1, "required_missing": []},
    }
    report = build_report_markdown(payload)
    assert "阻塞性缺口：1 项；提示性缺口：1 项" in report
    assert "未识别列：0 个；列冲突：1 个" in report
    assert "列名语义冲突" in report, "缺口类型应翻成中文"
    assert "阻塞 1 项 / 提示 1 项" in report


def test_p7_lists_unmapped_subjects(flow, allocation):
    payload = _payload(flow, allocation)
    payload["unmapped_subjects"] = [
        {"科目": "无形资产", "所属报表": "资产负债表", "金额（万元）": 4962.0,
         "登记状态": "未登记"},
        {"科目": "递延所得税资产", "所属报表": "资产负债表", "金额（万元）": 120.0,
         "登记状态": "已登记忽略"},
    ]
    p7 = build_report_markdown(payload).split("## P7 数据与计算质量审查")[1].split("## P8")[0]
    assert "未被任何规则关键词匹配的共 2 个（待登记 1 个）" in p7
    assert "递延所得税资产" in p7 and "已登记忽略" in p7


def test_report_marks_unimplemented_capabilities(flow, allocation):
    report = build_report_markdown(_payload(flow, allocation))
    assert report.count("统计预测未启用") >= 3, \
        "P1 情景与预测、P6 前瞻性风险、P8 第四节的措辞都要与实现一致"
    assert "跨期一致性比对" in report, "外部规则源监控已实现，不应再列为未启用"
    assert "情景模拟：本次未运行" in report, "未跑情景时如实说明"


def test_report_p9_appears_only_when_scenarios_ran(flow, allocation):
    payload = _payload(flow, allocation)
    assert "## P9" not in build_report_markdown(payload)

    payload["scenarios"] = [
        {"id": "base", "name": "基准", "spec": None, "changes": [],
         "rows": ROWS, "results": payload["results"],
         "allocation": allocation, "tax_flow": payload["tax_flow"], "errors": []},
        {"id": "qdmtt", "name": "QDMTT 落地",
         "spec": {"id": "qdmtt", "name": "QDMTT 落地",
                  "base": {"scenario_id": "default"},
                  "patch": [{"op": "set", "jurisdiction": "开曼群岛",
                             "field": "qdmtt_applies", "value": True}],
                  "assumptions": ["QDMTT 自 2025 年起适用"]},
         "changes": ["开曼群岛：QDMTT 适用 设为 是"],
         "rows": ROWS, "results": payload["results"],
         "allocation": allocation, "tax_flow": payload["tax_flow"], "errors": []},
    ]
    payload["scenario_comparison"] = {
        "base_id": "base", "base_name": "基准",
        "base_group": {"total_topup": 280.0, "qdmtt_retained": 0.0, "exported": 280.0,
                       "need_topup": 1},
        "scenarios": [{
            "id": "qdmtt", "name": "QDMTT 落地",
            "group": {"total_topup": {"base": 280.0, "target": 280.0, "delta": 0.0},
                      "qdmtt_retained": {"base": 0.0, "target": 280.0, "delta": 280.0},
                      "exported": {"base": 280.0, "target": 0.0, "delta": -280.0},
                      "need_topup": {"base": 1, "target": 1, "delta": 0}},
            "jurisdictions": [{
                "name": "开曼群岛",
                "etr": {"base": 0.01, "target": 0.01, "delta": 0.0},
                "topup_tax": {"base": 280.0, "target": 280.0, "delta": 0.0},
                "net_liability": {"base": 280.0, "target": 280.0, "delta": 0.0},
                "covered_taxes": {"base": 20.0, "target": 20.0, "delta": 0.0},
                "sbie": {"base": 0.0, "target": 0.0, "delta": 0.0},
                "risk": {"base": "high", "target": "high"},
            }, {
                "name": "德国",
                "etr": {"base": 0.26, "target": 0.26, "delta": 0.0},
                "topup_tax": {"base": 0.0, "target": 0.0, "delta": 0.0},
                # 应付款变了（QDMTT 让子公司自缴）→ 必须出现在差异表里
                "net_liability": {"base": 0.0, "target": 42.5, "delta": 42.5},
                "covered_taxes": {"base": 230.0, "target": 230.0, "delta": 0.0},
                "sbie": {"base": 0.0, "target": 0.0, "delta": 0.0},
                "risk": {"base": "safe_harbour", "target": "safe_harbour"},
            }],
        }],
    }
    payload["scenario_attribution"] = {
        "qdmtt": {"method": "sequential",
                  "order": ["分配开关（QDMTT / UTPR）"],
                  "factors": [{"factor": "分配开关（QDMTT / UTPR）",
                               "delta_topup": 0.0, "share": None}],
                  "total_delta": 0.0, "residual": 0.0},
    }
    report = build_report_markdown(payload)
    assert "## P9 情景对比" in report
    p9 = report.split("## P9 情景对比")[1].split("## 附录")[0]
    assert "QDMTT 落地" in p9 and "开曼群岛：QDMTT 适用 设为 是" in p9
    assert "QDMTT 自 2025 年起适用" in p9, "假设必须进报告"
    assert "QDMTT 留存" in p9 and "+280.00" in p9
    assert "逐步替换" in p9 and "残差" in p9
    assert "通常不改变集团补税总额" in p9, "必须写清 QDMTT 的真实影响"
    assert "德国" in p9 and "+42.50" in p9, "应付款变化必须出现在逐辖区差异里"
    assert "应付款 Δ" in p9
    assert "情景模拟：本次已运行 2 个情景" in report.split("## P2")[0]
    assert report.index("## P8 管理建议") < report.index("## P9 情景对比")


def test_report_p9_includes_cloud_narrative(flow, allocation):
    payload = _payload(flow, allocation)
    payload["scenarios"] = [{"id": "s1", "name": "情景A", "spec": None, "changes": [],
                             "rows": ROWS, "results": payload["results"],
                             "allocation": allocation, "tax_flow": payload["tax_flow"],
                             "errors": []}]
    payload["scenario_narrative"] = {
        "source": "llm",
        "analysis": "合并口径不变，但支付主体从母公司转到子公司。",
        "highlights": ["留存 0 → 280 万元"],
        "risks": ["假设未声明当地立法时点"],
        "followups": ["确认开曼 QDMTT 生效日"],
    }
    p9 = build_report_markdown(payload).split("## P9 情景对比")[1]
    assert "云端解读" in p9 and "来源：云端" in p9
    assert "支付主体从母公司转到子公司" in p9
    assert "确认开曼 QDMTT 生效日" in p9


def test_report_p9_flags_missing_assumptions(flow, allocation):
    payload = _payload(flow, allocation)
    payload["scenarios"] = [
        {"id": "s1", "name": "无假设情景",
         "spec": {"id": "s1", "name": "无假设情景", "base": {"scenario_id": "default"},
                  "patch": [{"op": "add", "jurisdiction": "德国", "field": "profit",
                             "value": 100}], "assumptions": []},
         "changes": ["德国：GloBE 利润 增加 100.00"], "rows": ROWS,
         "results": payload["results"], "allocation": allocation,
         "tax_flow": payload["tax_flow"], "errors": []},
    ]
    p9 = build_report_markdown(payload).split("## P9 情景对比")[1]
    assert "**未声明假设**" in p9


def test_report_without_cloud_data_still_renders(flow, allocation):
    results, tax_flow = flow
    report = build_report_markdown({"results": results, "rows": ROWS,
                                    "allocation": allocation, "tax_flow": tax_flow})
    assert "（本次运行未获得云端分析结论" in report
    assert "（无云端复核解读）" in report
    assert "（本次运行未获得云端建议动作" in report
    for page in PAGES:
        assert page in report, page


def test_report_without_results_is_safe():
    report = build_report_markdown({"results": [], "rows": []})
    assert "本轮没有可用的辖区计算结果" in report
    assert "## P1 执行摘要" not in report


def test_report_generated_at_can_be_injected(flow, allocation):
    payload = _payload(flow, allocation)
    payload["generated_at"] = "2026-09-26 12:00:00"
    assert "2026-09-26 12:00:00" in build_report_markdown(payload)
    payload.pop("generated_at")
    today = datetime.datetime.now().strftime("%Y-%m-%d")
    assert today in build_report_markdown(payload)


# ── 报告：QDMTT / IIR / UTPR 分配（P5）──

PARENT_CHILD_ROWS = [
    {"name": "母公司辖区", "profit": 1000.0, "current_tax": 200.0, "deferred_tax": 0.0,
     "revenue": 1000.0, "payroll": 100.0, "tangible_assets": 100.0, "parent_idx": None,
     "ownership": 1.0, "qdmtt_applies": False, "utpr_applies": True, "dtl_ledger": []},
    {"name": "低税子公司", "profit": 1000.0, "current_tax": 50.0, "deferred_tax": 0.0,
     "revenue": 1000.0, "payroll": 0.0, "tangible_assets": 0.0, "parent_idx": 0,
     "ownership": 1.0, "qdmtt_applies": False, "utpr_applies": True, "dtl_ledger": []},
]


def _scenario(rows):
    results = _assess(rows)
    alloc = run_allocation(
        results, {i: r.get("parent_idx") for i, r in enumerate(rows)},
        {i: True for i in range(len(rows))},
        qdmtt_applies={i: bool(r.get("qdmtt_applies")) for i, r in enumerate(rows)},
        ownership={i: float(r.get("ownership", 1.0)) for i, r in enumerate(rows)})
    tax_flow = compute_tax_flow(results, alloc, rows)
    return build_report_markdown({"results": results, "rows": rows,
                                  "allocation": alloc, "tax_flow": tax_flow})


def _with_qdmtt(rows, enabled: bool):
    rows = [dict(r) for r in rows]
    rows[1]["qdmtt_applies"] = enabled
    return rows


# 母公司持股 50%：子公司补税一半经 IIR 上收、一半进 UTPR 残池，
# 而子公司自己也会从残池分到钱 —— 此时"自身低税补缴"不能再计一次。
UTPR_ROWS = [
    {"name": "母公司辖区", "profit": 1000.0, "current_tax": 300.0, "deferred_tax": 0.0,
     "revenue": 1000.0, "payroll": 300.0, "tangible_assets": 300.0, "parent_idx": None,
     "ownership": 1.0, "qdmtt_applies": False, "utpr_applies": True, "dtl_ledger": []},
    {"name": "低税子公司", "profit": 1000.0, "current_tax": 50.0, "deferred_tax": 0.0,
     "revenue": 1000.0, "payroll": 100.0, "tangible_assets": 100.0, "parent_idx": 0,
     "ownership": 0.5, "qdmtt_applies": False, "utpr_applies": True, "dtl_ledger": []},
]


def test_p5_shows_iir_collection_and_net_liability():
    report = _scenario(_with_qdmtt(PARENT_CHILD_ROWS, False))
    p5 = report.split("## P5 QDMTT · IIR · UTPR 补税分配")[1].split("## P6")[0]
    assert "IIR 上收（母公司代缴）" in p5 and "100.00" in p5
    assert "低税子公司" in p5 and "代子公司缴（IIR）100.00" in p5
    assert "母公司辖区" in p5, "净负债表按支付主体列示"


def test_p6_flags_unprotected_jurisdiction_and_advises_qdmtt():
    report = _scenario(_with_qdmtt(PARENT_CHILD_ROWS, False))
    p6 = report.split("## P6 重点税务风险")[1].split("## P7")[0]
    assert "未实施 QDMTT，税源全部经 IIR / UTPR 流出" in p6
    p8 = report.split("## P8 管理建议")[1]
    assert "建议评估引入 QDMTT 的可行性" in p8


def test_p5_shows_qdmtt_retention_when_enacted():
    report = _scenario(_with_qdmtt(PARENT_CHILD_ROWS, True))
    p5 = report.split("## P5 QDMTT · IIR · UTPR 补税分配")[1].split("## P6")[0]
    assert "QDMTT 留存本国" in p5 and "QDMTT 自收 100.00" in p5
    p8 = report.split("## P8 管理建议")[1]
    assert "已通过 QDMTT 将 100.00 万元补税留在本国" in p8


def _net_liability_rows(report: str) -> list[tuple[float, list[float]]]:
    section = report.split("**最终补税净负债（按支付主体）**")[1]
    section = section.split("**UTPR 税源收入")[0]
    parsed = []
    for line in section.splitlines():
        if not line.startswith("|"):
            continue
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        if len(cells) != 3 or cells[0] == "支付主体":
            continue
        if not cells[1].replace(",", "").replace(".", "").isdigit():
            continue  # 表格分隔行
        total = float(cells[1].replace(",", ""))
        parts = [float(x.replace(",", ""))
                 for x in re.findall(r"[\d,]+\.\d{2}", cells[2])]
        parsed.append((total, parts))
    return parsed


def test_p5_net_liability_components_sum_to_total():
    """「构成」各项之和必须等于应付补税额：被 IIR 代缴的子公司补税不能算在自己头上。"""
    rows = _net_liability_rows(_scenario(UTPR_ROWS))
    assert rows, "净负债表应有内容"
    for total, parts in rows:
        assert parts, f"应付 {total} 万元缺少构成说明"
        assert sum(parts) == pytest.approx(total, abs=0.02), \
            f"构成 {parts} 之和与应付 {total} 万元不一致"


def _p5_overview_rows(report: str) -> dict[str, float]:
    """取 P5「分配总览」表：{项目: 金额}。"""
    section = report.split("**分配总览**")[1].split("**各辖区税源去向**")[0]
    out: dict[str, float] = {}
    for line in section.splitlines():
        if not line.startswith("|"):
            continue
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        if len(cells) != 2 or cells[0] in ("项目",):
            continue
        value = cells[1].replace(",", "")
        try:
            out[cells[0]] = float(value)
        except ValueError:
            continue
    return out


def test_report_p5_dispositions_sum_to_the_group_total():
    """P5 分配总览的四个去处（留存 / IIR / UTPR 分摊 / UPE 自身缴纳）必须加总到集团总额。

    踩过的坑：总览只列了前三项，相加比总额少 820.44（UPE 自己缴的那笔），
    读者一眼就能看出加不起来。ROWS 里开曼群岛是"自己有补缴的 UPE"，正好覆盖这条路径。
    """
    report = _scenario(ROWS)
    overview = _p5_overview_rows(report)
    assert overview, "P5 分配总览应有内容"
    total = overview["集团补税总额（按支付主体）"]
    upe_row = "UPE 自身缴纳（最终母公司直接缴纳）"
    assert upe_row in overview and overview[upe_row] > 0, \
        "该数据集必须包含「UPE 自己有补缴」的情形，否则这条检查是空过"
    dispositions = (overview.get("QDMTT 留存本国", 0.0)
                    + overview.get("IIR 上收（母公司代缴）", 0.0)
                    + overview.get("UTPR 已分摊", 0.0)
                    + overview[upe_row])
    assert dispositions == pytest.approx(total, abs=0.02), \
        f"四个去处相加 {dispositions} ≠ 集团总额 {total}"


def test_report_p1_sentence_names_the_third_destination():
    """P1 那句话必须讲清三个去处：留存 + 流出 + UPE 自身缴纳 = 补税合计。

    （只讲"留存 + 流出"时，两个数加不出总额 —— 真实数据里差 820.44。）
    """
    report = _scenario(ROWS)
    p1 = report.split("## P1")[1].split("## P2")[0]
    assert "QDMTT 留存" in p1 and "IIR / UTPR 流出" in p1
    assert "自身缴纳" in p1, "UPE 自己有补缴时，P1 必须写明第三个去处"
    numbers = [float(x.replace(",", ""))
               for x in re.findall(r"\*\*([\d,]+\.\d{2}) 万元\*\*", p1)]
    assert len(numbers) >= 4, f"应给出总额 + 三个去处，实际 {numbers}"
    assert numbers[1] + numbers[2] + numbers[3] == pytest.approx(numbers[0], abs=0.02), \
        f"三个去处 {numbers[1:4]} 相加 ≠ 补税合计 {numbers[0]}"


def test_report_numbers_match_engine_row_sums():
    """P2 集团汇总的合计必须等于引擎逐行加总（不能改成"利润 − SBIE"这种减法推导）。"""
    results = _assess(UTPR_ROWS)
    report = _scenario(UTPR_ROWS)
    section = report.split("**集团汇总**")[1].split("**辖区与集团架构**")[0]
    got: dict[str, float] = {}
    for line in section.splitlines():
        if not line.startswith("|"):
            continue
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        if len(cells) != 2 or cells[0] in ("项目",):
            continue
        try:
            got[cells[0]] = float(cells[1].replace(",", ""))
        except ValueError:
            continue
    expected = {
        "Covered Taxes 合计": round(sum(r["covered_taxes"] for r in results), 2),
        "SBIE 排除合计": round(sum(r["sbie"] for r in results), 2),
        "超额利润合计": round(sum(r["adjusted_profit"] for r in results), 2),
        "补税合计": round(sum((r["topup_tax"] or 0) for r in results), 2),
    }
    for label, value in expected.items():
        assert label in got, f"P2 集团汇总缺少「{label}」"
        assert got[label] == pytest.approx(value, abs=0.02), \
            f"{label}：报告 {got[label]} vs 引擎加总 {value}"


def test_report_p3_figures_match_engine_rows():
    """P3 逐辖区表里的利润/覆盖税额/ETR/超额利润/补税必须与引擎逐行一致。"""
    results = _assess(UTPR_ROWS)
    report = _scenario(UTPR_ROWS)
    section = report.split("## P3")[1].split("## P4")[0]
    parsed: dict[str, list[str]] = {}
    for line in section.splitlines():
        if not line.startswith("|"):
            continue
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        if len(cells) < 8 or cells[0] in ("辖区",) or set(cells[0]) <= {"-", ":"}:
            continue
        parsed[cells[0]] = cells
    assert parsed, "P3 应有辖区行"
    for row, item in zip(UTPR_ROWS, results):
        cells = parsed.get(row["name"])
        assert cells is not None, f"P3 缺少 {row['name']}"
        assert float(cells[1].replace(",", "")) == pytest.approx(
            round(item["profit"], 2), abs=0.02)
        assert float(cells[2].replace(",", "")) == pytest.approx(
            round(item["covered_taxes"], 2), abs=0.02)
        assert float(cells[5].replace(",", "")) == pytest.approx(
            round(item["adjusted_profit"], 2), abs=0.02)
        assert float(cells[7].replace(",", "")) == pytest.approx(
            round(item["topup_tax"] or 0, 2), abs=0.02)


def test_report_flags_upe_self_payment_instead_of_zero_retention():
    """各辖区税源去向表：UPE 那一行要显示"自身缴纳"，不能显示成留存率 0.0%。"""
    report = _scenario(ROWS)
    section = report.split("**各辖区税源去向**")[1].split("**最终补税净负债")[0]
    header = next(line for line in section.splitlines() if line.startswith("|"))
    assert "UPE 自身缴纳" in header, "税源去向表应有一列标明 UPE 自身缴纳"
    flagged = 0
    for line in section.splitlines():
        if not line.startswith("|"):
            continue
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        if len(cells) != 7 or cells[0] in ("辖区",) or set(cells[0]) <= {"-", ":"}:
            continue
        if cells[5] == "—":
            continue
        flagged += 1
        assert float(cells[5].replace(",", "")) > 0
        assert cells[6] == "—（自身缴纳）", \
            f"{cells[0]} 既有自身缴纳又标了留存率 {cells[6]}"
    assert flagged > 0, "该数据集必须包含 UPE 自己有补缴的情形，否则这条检查是空过"


# ── 报告：Word 版 ──

def _document_xml(data: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        return archive.read("word/document.xml").decode("utf-8")


def test_markdown_to_docx_returns_bytes_or_none(flow, allocation):
    report = build_report_markdown(_payload(flow, allocation))
    data = markdown_to_docx(report)
    if data is None:
        pytest.skip("环境未安装 python-docx")
    assert data[:2] == b"PK", "docx 应为 zip 容器"
    assert len(data) > 5000


def test_build_report_docx_matches_markdown_path(flow, allocation):
    data = build_report_docx(_payload(flow, allocation))
    if data is None:
        pytest.skip("环境未安装 python-docx")
    assert data[:2] == b"PK"


def test_docx_renders_markdown_tables_as_word_tables(flow, allocation):
    data = build_report_docx(_payload(flow, allocation))
    if data is None:
        pytest.skip("环境未安装 python-docx")
    xml = _document_xml(data)
    assert "<w:tbl>" in xml, "管道表格必须转成真正的 Word 表格"
    assert "| 辖区" not in xml, "不应把 Markdown 管道符原样写进 Word"


def test_docx_tables_fit_within_text_area(flow, allocation):
    """固定列宽后，每张表的合计宽度都不得超过正文宽度（否则会被撑出页面）。"""
    from docx import Document

    data = build_report_docx(_payload(flow, allocation))
    if data is None:
        pytest.skip("环境未安装 python-docx")
    document = Document(io.BytesIO(data))
    section = document.sections[0]
    available = section.page_width - section.left_margin - section.right_margin
    assert document.tables, "报告应包含表格"
    widest = max(len(table.columns) for table in document.tables)
    assert widest >= 9, "P3 辖区表应为 9 列（含判定）"
    for table in document.tables:
        total = sum(col.width for col in table.columns if col.width)
        assert total <= available, \
            f"{len(table.columns)} 列表格宽度 {total} 超出正文宽度 {available}"


def test_docx_breaks_pages_between_sections(flow, allocation):
    data = build_report_docx(_payload(flow, allocation))
    if data is None:
        pytest.skip("环境未安装 python-docx")
    xml = _document_xml(data)
    page_breaks = xml.count('w:type="page"')
    assert page_breaks >= len(PAGES), \
        f"P1–P8 应各占一页（分页符 {page_breaks} 个，不足 {len(PAGES)} 个）"

"""多辖区批量导入工具：解析辖区数据 + DTL 台账。"""

from __future__ import annotations

import io
from typing import Any

import openpyxl
import pandas as pd

from Agent.tools.base import run_tool
from utils import _parse_dataframe, parse_dtl_excel


def _read_upload_bytes(uploaded_file) -> bytes:
    if hasattr(uploaded_file, "getvalue"):
        data = uploaded_file.getvalue()
    else:
        data = uploaded_file.read()
        if hasattr(uploaded_file, "seek"):
            try:
                uploaded_file.seek(0)
            except Exception:
                pass
    if isinstance(data, str):
        data = data.encode("utf-8")
    return bytes(data)


def _find_sheet(sheet_names: list[str], keywords: tuple[str, ...],
                default: str) -> str:
    for name in sheet_names:
        if any(keyword in name for keyword in keywords):
            return name
    return default


def _parse_dtl_sheet(data: bytes, sheet_name: str) -> dict[str, Any]:
    """让 DTL 台账 sheet 成为 workbook 的 active sheet 后复用现有解析器。"""
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    if sheet_name not in wb.sheetnames:
        return {"dtl_by_jurisdiction": {}, "errors": []}
    for name in list(wb.sheetnames):
        if name != sheet_name:
            del wb[name]
    wb.active = 0
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return parse_dtl_excel(buf)


def _impl(uploaded_file, rates: dict | None = None) -> dict[str, Any]:
    data = _read_upload_bytes(uploaded_file)
    xls = pd.ExcelFile(io.BytesIO(data))
    sheet_names = list(xls.sheet_names)
    if not sheet_names:
        return {"rows": [], "errors": ["Excel 中没有 sheet"], "dtl_count": 0}

    main_sheet = _find_sheet(sheet_names, ("辖区数据", "辖区", "数据"), sheet_names[0])
    df = pd.read_excel(io.BytesIO(data), sheet_name=main_sheet)
    parsed = _parse_dataframe(df, rates)
    rows = parsed.get("rows", [])
    errors = list(parsed.get("errors", []))

    dtl_count = 0
    dtl_sheet = None
    for name in sheet_names:
        if "DTL" in name.upper() or "台账" in name:
            dtl_sheet = name
            break
    if dtl_sheet:
        dtl_result = _parse_dtl_sheet(data, dtl_sheet)
        errors.extend(dtl_result.get("errors", []))
        dtl_map = dtl_result.get("dtl_by_jurisdiction", {})
        for row in rows:
            row["dtl_ledger"] = dtl_map.get(row.get("name", ""), [])
            dtl_count += len(row["dtl_ledger"])

    return {
        "rows": rows,
        "errors": errors,
        "conversions": parsed.get("conversions", []),
        # 列级信号：未被识别的列、被排除的可疑映射、以及完整的列→字段对照。
        # 早期版本在这里把 unrecognized_columns 丢掉了，导致规则 Agent 看不到表头。
        "unrecognized_columns": parsed.get("unrecognized_columns", []),
        "column_conflicts": parsed.get("column_conflicts", []),
        "column_mapping": parsed.get("column_mapping", []),
        # 数据自己声明的金额单位（如「GloBE利润(万元)」）；没有则为空串，
        # 情景模拟会据此追问单位，避免十倍/百倍误差。
        "unit": parsed.get("unit", ""),
        "main_sheet": main_sheet,
        "dtl_sheet": dtl_sheet,
        "dtl_count": dtl_count,
    }


def parse_batch_workbook(uploaded_file, rates: dict | None = None,
                         run_id: str | None = None,
                         step: str | None = None) -> Any:
    """解析多辖区批量表，返回 ToolResult。"""
    return run_tool("parse_batch_workbook", _impl, uploaded_file,
                    rates=rates, run_id=run_id, step=step)

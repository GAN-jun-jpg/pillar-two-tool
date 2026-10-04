"""文件结构识别 Agent：判断上传 Excel 是多辖区批量表还是单体三大报表。"""

from __future__ import annotations

import io
import json
from typing import Any

import openpyxl

from Agent.agents.base import BaseAgent
from Agent.llm.prompts import SCHEMA_RECOGNITION_SYSTEM
from Agent.schemas import WorkflowState


def _read_upload_bytes(uploaded_file) -> bytes:
    if hasattr(uploaded_file, "getvalue"):
        data = uploaded_file.getvalue()
    else:
        data = uploaded_file.read()
        if hasattr(uploaded_file, "seek"):
            try:
                uploaded_file.seek(0)
            except Exception:
                pass
    if isinstance(data, str):
        data = data.encode("utf-8")
    return bytes(data)


class SchemaRecognitionAgent(BaseAgent):
    name = "schema_recognition"

    def _inspect_workbook(self, data: bytes) -> dict[str, Any]:
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
        sheet_names = list(wb.sheetnames)
        headers_by_sheet: dict[str, list[str]] = {}
        for name in sheet_names:
            ws = wb[name]
            rows = list(ws.iter_rows(min_row=1, max_row=3, values_only=True))
            header = []
            if rows:
                header = [str(v).strip() for v in rows[0] if v is not None]
            headers_by_sheet[name] = header
        wb.close()
        return {"sheet_names": sheet_names, "headers_by_sheet": headers_by_sheet}

    def _count_batch_rows(self, data: bytes, main_sheet: str) -> int:
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
        ws = wb[main_sheet]
        rows = list(ws.iter_rows(values_only=True))
        wb.close()
        if not rows:
            return 0
        header = [str(v).strip() if v is not None else "" for v in rows[0]]
        col_idx = None
        for idx, value in enumerate(header):
            if "辖区" in value or "国家" in value or "country" in value.lower():
                col_idx = idx
                break
        if col_idx is None:
            return max(0, len(rows) - 1)
        count = 0
        for row in rows[1:]:
            if col_idx < len(row) and row[col_idx] not in (None, ""):
                count += 1
        return count

    def _heuristic(self, info: dict[str, Any], data: bytes) -> dict[str, Any]:
        sheet_names = info["sheet_names"]
        headers_by_sheet = info["headers_by_sheet"]

        # 单体三大报表：sheet 名称或表头科目结构
        if any(any(kw in name for kw in ("利润表", "资产负债表", "现金流量表"))
               for name in sheet_names):
            return {"workbook_type": "financial_statements", "confidence": 0.95,
                    "reason": "检测到三大报表 sheet 名称", "source": "rule"}

        # 多辖区批量表：辖区 + 利润相关列
        for sheet_name, headers in headers_by_sheet.items():
            header_text = " ".join(headers)
            has_jurisdiction = any(("辖区" in h) or ("国家" in h) for h in headers)
            has_financial = any(("利润" in h) or ("GloBE" in h) or ("税额" in h)
                                for h in headers)
            if has_jurisdiction and has_financial:
                count = self._count_batch_rows(data, sheet_name)
                return {"workbook_type": "batch_jurisdictions", "confidence": 0.92,
                        "reason": f"检测到辖区/利润列，共 {count} 行", "source": "rule",
                        "main_sheet": sheet_name, "jurisdiction_count": count}

        # 单体财报常见科目列结构
        for sheet_name, headers in headers_by_sheet.items():
            header_text = " ".join(headers)
            if "科目" in header_text and not any("辖区" in h for h in headers):
                return {"workbook_type": "financial_statements", "confidence": 0.75,
                        "reason": "检测到科目列结构", "source": "rule"}

        return {"workbook_type": "unknown", "confidence": 0.2,
                "reason": "规则无法判断", "source": "rule"}

    def run(self, state: WorkflowState, uploaded_file: Any = None,
            **kwargs: Any) -> WorkflowState:
        state.set_step("schema_recognition")

        if uploaded_file is None:
            if state.parsed_data is not None:
                result = {
                    "workbook_type": "financial_statements",
                    "confidence": 1.0,
                    "reason": "输入为已解析的财报数据",
                    "source": "state",
                }
            else:
                result = {
                    "workbook_type": "unknown",
                    "confidence": 0.0,
                    "reason": "没有上传文件",
                    "source": "rule",
                }
        else:
            data = _read_upload_bytes(uploaded_file)
            info = self._inspect_workbook(data)
            result = self._heuristic(info, data)
            result.update(info)
            result["file_name"] = getattr(uploaded_file, "name", None)

            if result["workbook_type"] == "unknown" and self.llm_available():
                llm_result = self.llm_json(
                    SCHEMA_RECOGNITION_SYSTEM,
                    "请判断文件类型。只输出 JSON："
                    "{\"workbook_type\":\"batch_jurisdictions|financial_statements|unknown\","
                    "\"reason\":\"...\"}\n"
                    f"sheet_names={json.dumps(info['sheet_names'], ensure_ascii=False)}; "
                    f"headers_by_sheet={json.dumps(info['headers_by_sheet'], ensure_ascii=False)}",
                    fallback=None,
                )
                if isinstance(llm_result, dict) and llm_result.get("workbook_type") in (
                    "batch_jurisdictions", "financial_statements", "unknown"
                ):
                    result.update({
                        "workbook_type": llm_result["workbook_type"],
                        "reason": llm_result.get("reason", ""),
                        "source": "llm",
                        "confidence": 0.7,
                    })

        workbook_type = result.get("workbook_type")
        if workbook_type == "batch_jurisdictions":
            state.metadata["input_kind"] = "batch"
        elif workbook_type == "financial_statements":
            state.metadata["input_kind"] = "financial_statements"
        else:
            state.metadata["input_kind"] = "unknown"

        state.metadata["schema_recognition"] = result
        self.record(
            state,
            f"文件结构识别：{workbook_type}（{result.get('reason', '')}）",
            payload=result,
        )
        return state

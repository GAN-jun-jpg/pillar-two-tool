# -*- coding: utf-8 -*-
"""最小 Agent 集合测试。"""
import io

import openpyxl

from Agent.agents import (
    ChartAgent,
    DataAgent,
    PlannerAgent,
    ReviewAgent,
    SupervisorAgent,
    TaxAgent,
)
from Agent.schemas import WorkflowState


class Upload:
    def __init__(self, data: bytes, name: str):
        self._data = data
        self.name = name

    def read(self) -> bytes:
        return self._data


def _workbook_bytes() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "利润表"
    ws.append(["科目", "2025", "单位：万元"])
    ws.append(["营业总收入", 50000, ""])
    ws.append(["利润总额", 10000, ""])
    ws.append(["所得税费用", 1500, ""])
    ws.append(["递延所得税费用", 200, ""])

    ws2 = wb.create_sheet("资产负债表")
    ws2.append(["科目", "2025", "单位：万元"])
    ws2.append(["固定资产", 20000, ""])
    ws2.append(["资产总计", 50000, ""])

    ws3 = wb.create_sheet("现金流量表")
    ws3.append(["科目", "2025", "单位：万元"])
    ws3.append(["经营活动产生的现金流量净额", 8000, ""])
    ws3.append(["投资活动产生的现金流量净额", -2000, ""])
    ws3.append(["支付给职工以及为职工支付的现金", 1200, ""])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _row(**kw):
    base = {
        "name": "测试辖区", "profit": 1000.0, "current_tax": 100.0,
        "deferred_tax": 0.0, "revenue": 5000.0, "payroll": 0.0,
        "tangible_assets": 0.0, "parent_idx": None, "ownership": 1.0,
        "qdmtt_applies": False, "utpr_applies": True, "dtl_ledger": [],
    }
    base.update(kw)
    return base


def test_planner_agent_file_plan():
    state = WorkflowState.new()
    PlannerAgent().run(state, uploaded_file=Upload(b"x", "demo.xlsx"),
                       include_charts=True, export_gir=False)
    assert state.metadata["input_mode"] == "file"
    assert state.metadata["workflow_plan"] == [
        "parse", "map", "validate", "calculate", "chart"]


def test_data_agent_parse_and_map():
    state = WorkflowState.new()
    DataAgent().run(
        state,
        uploaded_file=Upload(_workbook_bytes(), "demo.xlsx"),
        jurisdiction_name="测试辖区",
    )
    assert state.mapped_rows[0]["payroll"] == 1200.0
    assert state.mapped_rows[0]["current_tax"] == 1300.0
    assert state.raw_rows == state.mapped_rows


def test_review_agent_validation_failure():
    state = WorkflowState.new()
    ReviewAgent().run(state, rows=[_row(name="")], calc_year=2024)
    assert state.status == "failed"
    assert state.validation_report["has_errors"] is True
    assert state.errors


def test_tax_agent_calculation():
    state = WorkflowState.new()
    TaxAgent().run(state, rows=[_row()], calc_year=2024)
    assert state.calculation_results[0]["topup_tax"] == 50.0
    assert "共 1 个辖区" in state.metadata["tax_summary"]


def test_chart_agent_generates_chart():
    state = WorkflowState.new()
    rows = [_row()]
    ReviewAgent().run(state, rows=rows, calc_year=2024)
    TaxAgent().run(state, rows=rows, calc_year=2024)
    ChartAgent(include_charts=True, export_gir=False).run(state, rows=rows)
    assert state.chart_data is not None
    assert state.chart_data["etr_bar"] is not None


def test_supervisor_agent_full_flow():
    state = SupervisorAgent(include_charts=False, export_gir=False).run(
        uploaded_file=Upload(_workbook_bytes(), "demo.xlsx"),
        jurisdiction_name="测试辖区",
        calc_year=2024,
    )
    assert state.status == "completed"
    assert state.mapped_rows[0]["payroll"] == 1200.0
    assert state.calculation_results[0]["current_tax"] == 1300.0
    assert state.messages
    assert state.tool_results


def test_supervisor_agent_missing_input():
    state = SupervisorAgent(include_charts=False, export_gir=False).run(calc_year=2024)
    assert state.status == "failed"
    assert state.errors

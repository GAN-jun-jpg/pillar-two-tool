# -*- coding: utf-8 -*-
"""Agent 线性编排器测试。"""
import io

import openpyxl

from Agent.orchestrator import WorkflowOrchestrator


class Upload:
    def __init__(self, data: bytes, name: str):
        self._data = data
        self.name = name

    def read(self) -> bytes:
        return self._data


def _workbook_bytes() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "利润表"
    ws.append(["科目", "2025", "单位：万元"])
    ws.append(["营业总收入", 50000, ""])
    ws.append(["利润总额", 10000, ""])
    ws.append(["所得税费用", 1500, ""])
    ws.append(["递延所得税费用", 200, ""])

    ws2 = wb.create_sheet("资产负债表")
    ws2.append(["科目", "2025", "单位：万元"])
    ws2.append(["固定资产", 20000, ""])
    ws2.append(["资产总计", 50000, ""])

    ws3 = wb.create_sheet("现金流量表")
    ws3.append(["科目", "2025", "单位：万元"])
    ws3.append(["经营活动产生的现金流量净额", 8000, ""])
    ws3.append(["投资活动产生的现金流量净额", -2000, ""])
    ws3.append(["支付给职工以及为职工支付的现金", 1200, ""])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _row(**kw):
    base = {
        "name": "测试辖区", "profit": 1000.0, "current_tax": 100.0,
        "deferred_tax": 0.0, "revenue": 5000.0, "payroll": 0.0,
        "tangible_assets": 0.0, "parent_idx": None, "ownership": 1.0,
        "qdmtt_applies": False, "utpr_applies": True, "dtl_ledger": [],
    }
    base.update(kw)
    return base


def test_orchestrator_full_file_flow():
    orchestrator = WorkflowOrchestrator(include_charts=True, export_gir=False)
    state = orchestrator.run(
        uploaded_file=Upload(_workbook_bytes(), "demo.xlsx"),
        jurisdiction_name="测试辖区",
        calc_year=2024,
    )
    assert state.status == "completed"
    assert state.parsed_data is not None
    assert state.mapped_rows[0]["payroll"] == 1200.0
    assert state.mapped_rows[0]["current_tax"] == 1300.0
    assert state.validation_report["has_errors"] is False
    assert state.calculation_results[0]["current_tax"] == 1300.0
    assert state.chart_data is not None
    assert state.chart_data["etr_bar"] is not None
    assert state.tool_results
    assert state.messages


def test_orchestrator_run_from_rows():
    orchestrator = WorkflowOrchestrator(include_charts=False, export_gir=False)
    state = orchestrator.run(rows=[_row()], calc_year=2024)
    assert state.status == "completed"
    assert state.calculation_results[0]["topup_tax"] == 50.0
    assert state.allocation is not None
    assert state.tax_flow is not None


def test_orchestrator_validation_failure_stops():
    orchestrator = WorkflowOrchestrator(include_charts=False, export_gir=False)
    state = orchestrator.run(rows=[_row(name="")], calc_year=2024)
    assert state.status == "failed"
    assert state.validation_report is not None
    assert state.validation_report["has_errors"] is True
    assert state.calculation_results == []


def test_orchestrator_missing_input():
    orchestrator = WorkflowOrchestrator(include_charts=False, export_gir=False)
    state = orchestrator.run(calc_year=2024)
    assert state.status == "failed"
    assert state.errors

def test_orchestrator_can_export_gir():
    orchestrator = WorkflowOrchestrator(include_charts=False, export_gir=True)
    state = orchestrator.run(rows=[_row()], calc_year=2024)
    assert state.status == "completed"
    assert state.export_info is not None
    assert state.export_info["gir_byte_length"] > 0

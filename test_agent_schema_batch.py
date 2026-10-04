# -*- coding: utf-8 -*-
"""文件结构识别和多辖区批量导入测试。"""
import io

import openpyxl

from Agent.agents import DataAgent, SchemaRecognitionAgent, SupervisorAgent
from Agent.schemas import WorkflowState
from Agent.tools import parse_batch_workbook


class Upload:
    def __init__(self, data: bytes, name: str):
        self._data = data
        self.name = name

    def read(self) -> bytes:
        return self._data

    def getvalue(self) -> bytes:
        return self._data


def _batch_bytes() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "辖区数据"
    ws.append(["辖区名称", "GloBE利润(万元)", "已缴税额(万元)", "递延所得税(万元)",
               "收入(万元)", "合格薪酬(万元)", "有形资产(万元)",
               "母公司", "持股比例", "QDMTT适用", "UTPR适用"])
    ws.append(["中国大陆", 1000, 100, 0, 5000, 50, 100, "", 1, "否", "是"])
    ws.append(["新加坡", 800, 80, 0, 4000, 40, 80, "中国大陆", 0.8, "是", "是"])

    ws2 = wb.create_sheet("DTL台账")
    ws2.append(["辖区", "DTL产生年份", "DTL产生金额(万元)", "DTL类型", "回转年份", "回转金额(万元)"])
    ws2.append(["中国大陆", 2020, 100, "固定资产", 2022, 60])
    ws2.append(["中国大陆", "", "", "", 2023, 40])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _financial_bytes() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "利润表"
    ws.append(["科目", "2025", "单位：万元"])
    ws.append(["营业总收入", 50000, ""])
    ws.append(["利润总额", 10000, ""])
    ws.append(["所得税费用", 1500, ""])
    ws.append(["递延所得税费用", 200, ""])
    ws2 = wb.create_sheet("资产负债表")
    ws2.append(["科目", "2025", "单位：万元"])
    ws2.append(["固定资产", 20000, ""])
    ws2.append(["资产总计", 50000, ""])
    ws3 = wb.create_sheet("现金流量表")
    ws3.append(["科目", "2025", "单位：万元"])
    ws3.append(["经营活动产生的现金流量净额", 8000, ""])
    ws3.append(["投资活动产生的现金流量净额", -2000, ""])
    ws3.append(["支付给职工以及为职工支付的现金", 1200, ""])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_schema_recognition_batch():
    state = WorkflowState.new()
    SchemaRecognitionAgent().run(state, uploaded_file=Upload(_batch_bytes(), "batch.xlsx"))
    info = state.metadata["schema_recognition"]
    assert info["workbook_type"] == "batch_jurisdictions"
    assert info["jurisdiction_count"] == 2
    assert state.metadata["input_kind"] == "batch"


def test_schema_recognition_financial_statements():
    state = WorkflowState.new()
    SchemaRecognitionAgent().run(state, uploaded_file=Upload(_financial_bytes(), "fs.xlsx"))
    info = state.metadata["schema_recognition"]
    assert info["workbook_type"] == "financial_statements"
    assert state.metadata["input_kind"] == "financial_statements"


def test_batch_import_tool_parses_rows_and_dtl():
    result = parse_batch_workbook(Upload(_batch_bytes(), "batch.xlsx"))
    assert result.ok is True
    data = result.data
    assert len(data["rows"]) == 2
    assert data["dtl_count"] == 1
    assert data["rows"][0]["name"] == "中国大陆"
    assert data["rows"][0]["dtl_ledger"][0]["amount"] == 100


def test_data_agent_batch_path():
    state = WorkflowState.new()
    state.metadata["input_kind"] = "batch"
    DataAgent().run(state, uploaded_file=Upload(_batch_bytes(), "batch.xlsx"))
    assert len(state.mapped_rows) == 2
    assert state.mapped_rows[0]["name"] == "中国大陆"
    assert state.metadata["batch_import"]["dtl_count"] == 1


def test_supervisor_batch_flow():
    state = SupervisorAgent(include_charts=False, export_gir=False).run(
        uploaded_file=Upload(_batch_bytes(), "batch.xlsx"),
        calc_year=2024,
    )
    assert state.status == "completed"
    assert len(state.mapped_rows) == 2
    tool_names = [r.tool_name for r in state.tool_results]
    assert "parse_batch_workbook" in tool_names
    assert "parse_financial_statement" not in tool_names

"""GIR (GloBE Information Return) 标准格式 Excel 导出 v3.0。

基于 OECD GloBE Model Rules 申报要求，将 Pillar Two 计算结果
生成为事务所级格式化 Excel 工作簿（5 Sheets + 图表嵌入 + 合规追溯矩阵）。
"""

from __future__ import annotations

import io
from datetime import datetime

import openpyxl
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import (
    Alignment, Border, Font, PatternFill, Side,
)
from openpyxl.utils import get_column_letter

import plotly.io as pio

from calculator import build_dtl_schedule, get_dtl_status, DTL_TYPES, get_sbie_rates, SIMPLIFIED_ETR_THRESHOLDS

# ═══════════════════════════════════════════════════════════
# 样式常量
# ═══════════════════════════════════════════════════════════
HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
HEADER_FONT = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
SUBHEADER_FILL = PatternFill(start_color="D6E4F0", end_color="D6E4F0", fill_type="solid")
TITLE_FONT = Font(name="微软雅黑", size=16, bold=True, color="1F4E79")
SECTION_FONT = Font(name="微软雅黑", size=12, bold=True, color="1F4E79")
SUBTITLE_FONT = Font(name="微软雅黑", size=11, color="555555")
BODY_FONT = Font(name="微软雅黑", size=10)
BOLD_FONT = Font(name="微软雅黑", size=10, bold=True)
SMALL_FONT = Font(name="微软雅黑", size=9, color="888888")
NUMBER_FONT = Font(name="Consolas", size=10)
TOTAL_FILL = PatternFill(start_color="E8ECF1", end_color="E8ECF1", fill_type="solid")

THIN_BORDER = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
    bottom=Side(style="thin", color="D9D9D9"),
)

CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
RIGHT = Alignment(horizontal="right", vertical="center")

RISK_FILLS = {
    "safe_harbour": PatternFill(start_color="D4EDDA", end_color="D4EDDA", fill_type="solid"),
    "low": PatternFill(start_color="E8F5E9", end_color="E8F5E9", fill_type="solid"),
    "high": PatternFill(start_color="FFE8E8", end_color="FFE8E8", fill_type="solid"),
    "n/a": PatternFill(start_color="F5F5F5", end_color="F5F5F5", fill_type="solid"),
}

RISK_LABELS = {
    "safe_harbour": "🛡️ Safe Harbour",
    "low": "✅ 安全",
    "high": "⚠️ 需补税",
    "n/a": "—",
}

SH_ICONS = {
    "De Minimis": "📏",
    "Simplified ETR": "📊",
    "Routine Profit": "🏭",
}

TRACE_STATUS_ICONS = {
    "applied": "✅ 已适用",
    "triggered": "⚠️ 已触发",
    "exempt": "🛡️ 豁免",
    "not_applicable": "— 不适用",
}

# ═══════════════════════════════════════════════════════════
# 样式工具函数
# ═══════════════════════════════════════════════════════════

def _style_header(ws, row: int, col_count: int) -> None:
    for c in range(1, col_count + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = THIN_BORDER


def _style_body(ws, row: int, col_count: int, font: Font | None = None) -> None:
    for c in range(1, col_count + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = font or BODY_FONT
        cell.alignment = CENTER if c > 1 else LEFT
        cell.border = THIN_BORDER


def _style_total_row(ws, row: int, col_count: int) -> None:
    for c in range(1, col_count + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = BOLD_FONT
        cell.fill = TOTAL_FILL
        cell.border = Border(
            top=Side(style="medium", color="1F4E79"),
            bottom=Side(style="medium", color="1F4E79"),
            left=Side(style="thin", color="D9D9D9"),
            right=Side(style="thin", color="D9D9D9"),
        )
        cell.alignment = CENTER if c > 1 else LEFT


def _auto_width(ws, col_count: int, min_width: int = 10, max_width: int = 42) -> None:
    for c in range(1, col_count + 1):
        letter = get_column_letter(c)
        max_len = min_width
        for row in ws.iter_rows(min_col=c, max_col=c, values_only=True):
            for val in row:
                if val is not None:
                    max_len = max(max_len, min(len(str(val)) * 1.3 + 2, max_width))
        ws.column_dimensions[letter].width = max_len


def _setup_print(ws, orientation: str = "landscape") -> None:
    ws.sheet_properties.pageSetUpPr = openpyxl.worksheet.properties.PageSetupProperties(fitToPage=True)
    ws.page_setup.orientation = orientation
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.freeze_panes = "A2"



def _style_for_excel(fig):
    """导出前把深色主题图表转为白底深字版本，保证在 Excel 中可读。"""
    import copy
    fig = copy.deepcopy(fig)
    fig.update_layout(
        paper_bgcolor="white",
        plot_bgcolor="white",
        font=dict(color="#1F2937", family="Microsoft YaHei"),
        title_font_color="#1F4E79",
    )
    fig.update_xaxes(gridcolor="#E5E7EB", tickfont=dict(color="#374151"))
    fig.update_yaxes(gridcolor="#E5E7EB", tickfont=dict(color="#374151"))
    for tr in fig.data:
        if tr.type == "sankey":
            tr.textfont = dict(color="#111827", size=11, family="Microsoft YaHei")
    for ann in (fig.layout.annotations or []):
        ann.bgcolor = "rgba(255,255,255,0.8)"
        if ann.font:
            ann.font.color = "#374151"
    return fig

def _fig_to_excel_image(fig, width: int = 1100, height: int = 360) -> XLImage | None:
    """将 Plotly figure 转为 openpyxl 可嵌入的图片对象。"""
    if fig is None:
        return None
    try:
        img_bytes = pio.to_image(_style_for_excel(fig), format="png", width=width, height=height, scale=2)
        img = XLImage(io.BytesIO(img_bytes))
        img.width = width * 0.75
        img.height = height * 0.75
        return img
    except Exception:
        return None


def _next_row(ws) -> int:
    return ws.max_row + 1


# ═══════════════════════════════════════════════════════════
# Sheet 1: 封面与摘要（合并）
# ═══════════════════════════════════════════════════════════

def _write_cover_summary(
    ws,
    group_name: str,
    sbie_year: int,
    results: list[dict],
    rows: list[dict],
    allocation: dict,
    tax_flow: dict | None,
    total_topup: float,
    sh_count: int,
    high_count: int,
    total_retained: float,
    total_exported: float,
    etr_bar_fig,  # Plotly figure | None
) -> None:
    """封面 — 基本信息 + 关键指标 + ETR 概览图。"""
    _pr, _ar = get_sbie_rates(sbie_year)
    _is_transition = sbie_year <= 2032

    # ── 标题区 ──
    ws.merge_cells("A1:H1")
    ws["A1"] = "GloBE Information Return (GIR)"
    ws["A1"].font = TITLE_FONT
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 36

    ws.merge_cells("A2:H2")
    ws["A2"] = "Pillar Two 全球最低税负 — 标准申报信息汇总"
    ws["A2"].font = SUBTITLE_FONT
    ws["A2"].alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells("A3:H3")
    ws["A3"] = f"导出时间：{datetime.now().strftime('%Y-%m-%d %H:%M')}　　工具版本：Pillar Two GIR v3.0"
    ws["A3"].font = SMALL_FONT
    ws["A3"].alignment = Alignment(horizontal="center", vertical="center")

    # ── 左栏：基本信息 ──
    info = [
        ("集团名称", group_name),
        ("申报财年", f"{sbie_year} 年"),
        ("辖区数量", f"{len(results)} 个"),
        ("SBIE 排除率", f"薪酬 {_pr:.1%} / 资产 {_ar:.1%}" + ("（过渡期）" if _is_transition else "（稳态）")),
        ("计算规则", "OECD GloBE Model Rules · QDMTT → IIR → UTPR"),
    ]
    for i, (label, value) in enumerate(info):
        r = 5 + i
        ws.cell(row=r, column=1, value=label).font = BOLD_FONT
        ws.cell(row=r, column=1).alignment = RIGHT
        ws.merge_cells(f"B{r}:D{r}")
        ws.cell(row=r, column=2, value=str(value)).font = BODY_FONT
        ws.cell(row=r, column=2).alignment = LEFT

    # ── 右栏：关键指标 ──
    metrics = [
        ("集团补税总额（万元）", f"{total_topup:,.2f}", "⚠️" if total_topup > 0 else "✅"),
        ("Safe Harbour 豁免", f"{sh_count} 个辖区", "🛡️" if sh_count > 0 else ""),
        ("需补税辖区", f"{high_count} 个", "⚠️" if high_count > 0 else "✅"),
        ("QDMTT 自留（万元）", f"{total_retained:,.2f}", "🏛️"),
        ("税源流出（万元）", f"{total_exported:,.2f}", "📤"),
        ("ETR ≥ 15% 辖区", f"{sum(1 for r in results if r.get('risk') == 'low')} 个", "✅"),
    ]
    for j, (label, value, icon) in enumerate(metrics):
        r = 5 + j
        ws.cell(row=r, column=5, value=f"{icon} {label}").font = BOLD_FONT
        ws.cell(row=r, column=5).alignment = RIGHT
        ws.merge_cells(f"F{r}:H{r}")
        ws.cell(row=r, column=6, value=value).font = NUMBER_FONT
        ws.cell(row=r, column=6).alignment = LEFT

    # ── 辖区风险一览 ──
    tbl_start = 5 + max(len(info), len(metrics)) + 1
    ws.merge_cells(f"A{tbl_start}:H{tbl_start}")
    ws.cell(row=tbl_start, column=1, value="── 辖区风险一览 ──").font = BOLD_FONT

    hdrs = ["#", "辖区", "GloBE 利润（万）", "ETR", "风险等级", "补税（万）", "Safe Harbour", "处理方式"]
    hr = tbl_start + 1
    for c, h in enumerate(hdrs, 1):
        ws.cell(row=hr, column=c, value=h)
    _style_header(ws, hr, len(hdrs))

    for i, (r, row_data) in enumerate(zip(results, rows)):
        rn = hr + 1 + i
        topup = r.get("topup_tax") or 0.0
        sh_rule = r.get("safe_harbour")
        qdmtt = row_data.get("qdmtt_applies", False)
        parent = row_data.get("parent_idx")

        if sh_rule:
            handling = f"🛡️ {sh_rule}"
        elif topup > 0 and qdmtt:
            handling = "🏛️ QDMTT 自收"
        elif topup > 0 and parent is not None:
            handling = f"🔗 IIR → {rows[parent].get('name', '?')}"
        elif topup > 0:
            handling = "📐 UTPR 分摊"
        else:
            handling = "—"

        vals = [
            i + 1,
            row_data.get("name", ""),
            r.get("profit", 0),
            r.get("etr"),
            RISK_LABELS.get(r.get("risk", ""), ""),
            topup,
            sh_rule or "—",
            handling,
        ]
        for c, v in enumerate(vals, 1):
            ws.cell(row=rn, column=c, value=v)
        _style_body(ws, rn, len(hdrs))

        ws.cell(row=rn, column=3).number_format = '#,##0.00'
        ws.cell(row=rn, column=3).font = NUMBER_FONT
        ws.cell(row=rn, column=3).alignment = RIGHT
        if isinstance(vals[3], (int, float)):
            ws.cell(row=rn, column=4).number_format = '0.00%'
            ws.cell(row=rn, column=4).font = NUMBER_FONT
            ws.cell(row=rn, column=4).alignment = RIGHT
        ws.cell(row=rn, column=6).number_format = '#,##0.00'
        ws.cell(row=rn, column=6).font = NUMBER_FONT
        ws.cell(row=rn, column=6).alignment = RIGHT

        risk_fill = RISK_FILLS.get(r.get("risk", ""))
        if risk_fill:
            ws.cell(row=rn, column=5).fill = risk_fill

    # ── ETR 概览图嵌入 ──
    chart_row = hr + 1 + len(results) + 2
    etr_img = _fig_to_excel_image(etr_bar_fig, width=1100, height=340)
    if etr_img:
        ws.merge_cells(f"A{chart_row}:H{chart_row}")
        ws.cell(row=chart_row, column=1, value="── 各辖区 ETR 对比 ──").font = BOLD_FONT
        ws.add_image(etr_img, f"A{chart_row + 1}")

    # ── 免责声明 ──
    disc_row = chart_row + 18 if etr_img else chart_row + 2
    ws.merge_cells(f"A{disc_row}:H{disc_row}")
    ws.cell(row=disc_row, column=1,
            value="⚠️ 免责声明：本工具为 Pillar Two 规则学习与模拟演示用途，不构成税务意见。"
                  "实际申报请以 OECD 最新指引及当地税务机关要求为准。").font = SMALL_FONT
    ws.cell(row=disc_row, column=1).alignment = LEFT

    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 18
    ws.column_dimensions["D"].width = 12
    ws.column_dimensions["E"].width = 20
    ws.column_dimensions["F"].width = 18
    ws.column_dimensions["G"].width = 18
    ws.column_dimensions["H"].width = 20


# ═══════════════════════════════════════════════════════════
# Sheet 2: ETR 计算与补税分配（合并）
# ═══════════════════════════════════════════════════════════

def _write_etr_and_allocation(
    ws,
    results: list[dict],
    rows: list[dict],
    allocation: dict,
    waterfall_fig,  # Plotly figure | None
) -> None:
    """ETR 计算明细 + 补税分配 + 瀑布图嵌入。"""
    # ── Part A: ETR 计算明细 ──
    ws.merge_cells("A1:K1")
    ws["A1"] = "Part A — ETR 计算明细"
    ws["A1"].font = SECTION_FONT

    headers = [
        "辖区", "GloBE 利润", "SBIE 排除", "超额利润",
        "当期所得税", "递延所得税", "DTL 回转惩罚", "Covered Taxes",
        "ETR", "风险等级", "补税金额",
    ]
    for c, h in enumerate(headers, 1):
        ws.cell(row=3, column=c, value=h)
    _style_header(ws, 3, len(headers))

    num_cols = (2, 3, 4, 5, 6, 7, 8, 11)
    totals = {c: 0.0 for c in num_cols}

    for i, (r, row_data) in enumerate(zip(results, rows)):
        rn = 4 + i
        vals = [
            row_data.get("name", ""),
            r.get("profit", 0),
            r.get("sbie", 0),
            r.get("adjusted_profit", 0),
            r.get("current_tax", 0),
            r.get("deferred_tax", 0),
            r.get("recapture_amount", 0),
            r.get("covered_taxes", 0),
            r.get("etr"),
            RISK_LABELS.get(r.get("risk", ""), ""),
            r.get("topup_tax") or 0.0,
        ]
        for c, v in enumerate(vals, 1):
            ws.cell(row=rn, column=c, value=v)
        _style_body(ws, rn, len(headers))

        for c in num_cols:
            cell = ws.cell(row=rn, column=c)
            cell.number_format = '#,##0.00'
            cell.font = NUMBER_FONT
            cell.alignment = RIGHT
            if isinstance(vals[c - 1], (int, float)):
                totals[c] += vals[c - 1]

        etr_cell = ws.cell(row=rn, column=9)
        if isinstance(vals[8], (int, float)):
            etr_cell.number_format = '0.00%'
        etr_cell.font = NUMBER_FONT
        etr_cell.alignment = RIGHT

        risk_fill = RISK_FILLS.get(r.get("risk", ""))
        if risk_fill:
            ws.cell(row=rn, column=10).fill = risk_fill

    # 合计行
    tr = 4 + len(results)
    ws.cell(row=tr, column=1, value="合计").font = BOLD_FONT
    for ci in range(2, len(headers) + 1):
        cell = ws.cell(row=tr, column=ci)
        if ci in num_cols:
            cell.value = totals[ci]
            cell.number_format = '#,##0.00'
            cell.font = NUMBER_FONT
        elif ci == 9:
            cell.value = "—"
            cell.font = BOLD_FONT
        else:
            cell.font = BOLD_FONT
    _style_total_row(ws, tr, len(headers))

    # ── 瀑布图嵌入 ──
    wf_row = tr + 3
    wf_img = _fig_to_excel_image(waterfall_fig, width=1100, height=380)
    if wf_img:
        ws.merge_cells(f"A{wf_row}:K{wf_row}")
        ws.cell(row=wf_row, column=1, value="── Covered Taxes 构成（万元）──").font = BOLD_FONT
        ws.add_image(wf_img, f"A{wf_row + 1}")
        alloc_start = wf_row + 22
    else:
        alloc_start = tr + 3

    # ── Part B: 补税分配 ──
    ws.merge_cells(f"A{alloc_start}:G{alloc_start}")
    ws.cell(row=alloc_start, column=1, value="Part B — 补税分配（QDMTT → IIR → UTPR）").font = SECTION_FONT

    alloc_headers = ["辖区", "QDMTT 自收", "IIR 代缴（收）", "IIR 代缴（支）", "UTPR 分摊", "净负债"]
    ah_row = alloc_start + 1
    for c, h in enumerate(alloc_headers, 1):
        ws.cell(row=ah_row, column=c, value=h)
    _style_header(ws, ah_row, len(alloc_headers))

    qdmtt_collected = allocation["qdmtt"]["qdmtt_collected"]
    iir_collected = allocation["iir"]["collected"]
    iir_flows = allocation["iir"]["flows"]
    utpr_alloc = allocation["utpr"]["allocated"]
    net = allocation["net_liability"]

    iir_paid: dict[int, float] = {}
    for f in iir_flows:
        child = f["from"]
        iir_paid[child] = iir_paid.get(child, 0.0) + f["amount"]

    t_qdmtt = t_iir_in = t_iir_out = t_utpr = t_net = 0.0
    for i, row in enumerate(rows):
        name = row.get("name", "").strip()
        if not name:
            continue
        rn = ah_row + 1 + i
        qdmtt_amt = qdmtt_collected.get(i, 0.0)
        iir_in = iir_collected.get(i, 0.0)
        iir_out = iir_paid.get(i, 0.0)
        utpr_amt = utpr_alloc.get(i, 0.0)
        net_amt = net.get(i, 0.0)

        t_qdmtt += qdmtt_amt
        t_iir_in += iir_in
        t_iir_out += iir_out
        t_utpr += utpr_amt
        t_net += net_amt

        vals = [name, qdmtt_amt, iir_in, iir_out, utpr_amt, net_amt]
        for c, v in enumerate(vals, 1):
            ws.cell(row=rn, column=c, value=v)
        _style_body(ws, rn, len(alloc_headers))
        for c in (2, 3, 4, 5, 6):
            cell = ws.cell(row=rn, column=c)
            cell.number_format = '#,##0.00'
            cell.font = NUMBER_FONT
            cell.alignment = RIGHT

    tr2 = ah_row + 1 + len(rows)
    total_vals = ["合计", t_qdmtt, t_iir_in, t_iir_out, t_utpr, t_net]
    for c, v in enumerate(total_vals, 1):
        ws.cell(row=tr2, column=c, value=v if v != "" else None)
    _style_total_row(ws, tr2, len(alloc_headers))
    for c in (2, 3, 4, 5, 6):
        ws.cell(row=tr2, column=c).number_format = '#,##0.00'
        ws.cell(row=tr2, column=c).font = NUMBER_FONT

    _auto_width(ws, max(len(headers), len(alloc_headers)), max_width=42)
    _setup_print(ws)


# ═══════════════════════════════════════════════════════════
# Sheet 3: Safe Harbour 与税源流向（合并）
# ═══════════════════════════════════════════════════════════

def _write_sh_and_taxflow(
    ws,
    results: list[dict],
    rows: list[dict],
    sbie_year: int,
    tax_flow: dict | None,
    sankey_fig,  # Plotly figure | None
) -> None:
    """Safe Harbour 明细 + 税源流向 + Sankey 图嵌入。"""
    # ── Part A: Safe Harbour 明细 ──
    ws.merge_cells("A1:G1")
    ws["A1"] = "Part A — Safe Harbour 安全港豁免明细"
    ws["A1"].font = SECTION_FONT

    sh_headers = ["辖区", "适用规则", "收入（万）", "GloBE 利润（万）", "SBIE（万）", "ETR", "判定说明"]
    hr = 3
    for c, h in enumerate(sh_headers, 1):
        ws.cell(row=hr, column=c, value=h)
    _style_header(ws, hr, len(sh_headers))

    sh_rn = hr + 1
    for r, row_data in zip(results, rows):
        sh_rule = r.get("safe_harbour")
        if not sh_rule:
            continue
        name = row_data.get("name", "")
        revenue = row_data.get("revenue", 0)
        profit = r.get("profit", 0)
        sbie = r.get("sbie", 0)
        etr = r.get("etr")

        if sh_rule == "De Minimis":
            explanation = "收入 ≤ €10M 且 利润 ≤ €1M → 满足 De Minimis 测试"
        elif sh_rule == "Simplified ETR":
            threshold = SIMPLIFIED_ETR_THRESHOLDS.get(sbie_year, 0.15)
            explanation = f"ETR {etr:.1%} ≥ {threshold:.1%}（{sbie_year} 年过渡期阈值）"
        elif sh_rule == "Routine Profit":
            explanation = f"GloBE 利润 {profit:,.0f}万 ≤ SBIE {sbie:,.0f}万"
        else:
            explanation = sh_rule

        vals = [name, f"{SH_ICONS.get(sh_rule, '')} {sh_rule}", revenue, profit, sbie, etr, explanation]
        for c, v in enumerate(vals, 1):
            ws.cell(row=sh_rn, column=c, value=v)
        _style_body(ws, sh_rn, len(sh_headers))
        for c in (3, 4, 5):
            cell = ws.cell(row=sh_rn, column=c)
            cell.number_format = '#,##0.00'
            cell.font = NUMBER_FONT
            cell.alignment = RIGHT
        if isinstance(etr, (int, float)):
            cell = ws.cell(row=sh_rn, column=6)
            cell.number_format = '0.00%'
            cell.font = NUMBER_FONT
            cell.alignment = RIGHT
        for c in range(1, len(sh_headers) + 1):
            ws.cell(row=sh_rn, column=c).fill = PatternFill(
                start_color="D4EDDA", end_color="D4EDDA", fill_type="solid")
        sh_rn += 1

    if sh_rn == hr + 1:
        ws.merge_cells(f"A{sh_rn}:G{sh_rn}")
        ws.cell(row=sh_rn, column=1, value="无辖区通过 Safe Harbour 测试。").font = SUBTITLE_FONT
        ws.cell(row=sh_rn, column=1).alignment = CENTER
        sh_rn += 1

    # ── Part B: 税源流向 ──
    tf_start = sh_rn + 2
    ws.merge_cells(f"A{tf_start}:G{tf_start}")
    ws.cell(row=tf_start, column=1, value="Part B — 税源流向分析").font = SECTION_FONT

    if tax_flow and tax_flow.get("sources"):
        tf_headers = ["#", "辖区", "补税（万）", "QDMTT 自留（万）", "IIR 流出（万）",
                       "UTPR 流出（万）", "自留率"]
        hr2 = tf_start + 1
        for c, h in enumerate(tf_headers, 1):
            ws.cell(row=hr2, column=c, value=h)
        _style_header(ws, hr2, len(tf_headers))

        for i, s in enumerate(tax_flow["sources"]):
            rn = hr2 + 1 + i
            retention = s.get("retention_rate")
            retention_str = f"{retention:.1%}" if retention is not None else "—"
            vals = [i + 1, s["name"], s["topup"], s["qdmtt_retained"],
                    s["iir_exported"], s["utpr_exported"], retention_str]
            for c, v in enumerate(vals, 1):
                ws.cell(row=rn, column=c, value=v)
            _style_body(ws, rn, len(tf_headers))
            for c in (3, 4, 5, 6):
                cell = ws.cell(row=rn, column=c)
                cell.number_format = '#,##0.00'
                cell.font = NUMBER_FONT
                cell.alignment = RIGHT
            if s.get("pure_export"):
                for c in range(1, len(tf_headers) + 1):
                    ws.cell(row=rn, column=c).fill = PatternFill(
                        start_color="FFF3CD", end_color="FFF3CD", fill_type="solid")

        tf_end = hr2 + 1 + len(tax_flow["sources"])
    else:
        ws.merge_cells(f"A{tf_start+1}:G{tf_start+1}")
        ws.cell(row=tf_start + 1, column=1, value="无低税辖区。").font = SUBTITLE_FONT
        tf_end = tf_start + 2

    # ── 接收方 ──
    if tax_flow and tax_flow.get("utpr_recipients"):
        rec_start = tf_end + 2
        ws.merge_cells(f"A{rec_start}:G{rec_start}")
        ws.cell(row=rec_start, column=1, value="UTPR 接收方").font = SECTION_FONT

        rec_hdrs = ["#", "辖区", "UTPR 接收（万）", "占残池份额"]
        rh_row = rec_start + 1
        for c, h in enumerate(rec_hdrs, 1):
            ws.cell(row=rh_row, column=c, value=h)
        _style_header(ws, rh_row, len(rec_hdrs))
        for i, rec in enumerate(tax_flow["utpr_recipients"]):
            rn = rh_row + 1 + i
            vals = [i + 1, rec["name"], rec["utpr_received"], rec.get("share_of_pool", 0)]
            for c, v in enumerate(vals, 1):
                ws.cell(row=rn, column=c, value=v)
            _style_body(ws, rn, len(rec_hdrs))
            ws.cell(row=rn, column=3).number_format = '#,##0.00'
            ws.cell(row=rn, column=3).font = NUMBER_FONT
            ws.cell(row=rn, column=3).alignment = RIGHT
            ws.cell(row=rn, column=4).number_format = '0.0%'
            ws.cell(row=rn, column=4).font = NUMBER_FONT
            ws.cell(row=rn, column=4).alignment = RIGHT
        chart_anchor = rh_row + 1 + len(tax_flow["utpr_recipients"]) + 2
    else:
        chart_anchor = tf_end + 2

    # ── Sankey 图嵌入 ──
    sankey_img = _fig_to_excel_image(sankey_fig, width=1100, height=400)
    if sankey_img:
        ws.merge_cells(f"A{chart_anchor}:G{chart_anchor}")
        ws.cell(row=chart_anchor, column=1, value="── 税源流向全景图 ──").font = BOLD_FONT
        ws.add_image(sankey_img, f"A{chart_anchor + 1}")

    _auto_width(ws, 7, max_width=55)
    _setup_print(ws)


# ═══════════════════════════════════════════════════════════
# Sheet 4: DTL 台账
# ═══════════════════════════════════════════════════════════

def _write_dtl(ws, rows: list[dict], calc_year: int) -> None:
    headers = [
        "辖区", "DTL ID", "产生年份", "原始金额", "类型",
        "Qualified", "已回转", "剩余未回转", "状态", "到期年份",
        "Recapture 金额", "时间轴",
    ]
    for c, h in enumerate(headers, 1):
        ws.cell(row=1, column=c, value=h)
    _style_header(ws, 1, len(headers))

    rn = 2
    jurisdiction_dtl_totals: dict[str, dict] = {}

    for row in rows:
        name = row.get("name", "").strip()
        if not name:
            continue
        if name not in jurisdiction_dtl_totals:
            jurisdiction_dtl_totals[name] = {"original": 0.0, "reversed": 0.0,
                                             "remaining": 0.0, "recaptured": 0.0}

        jur_start_rn = rn
        for entry in row.get("dtl_ledger", []):
            dtl_id = entry.get("id", "")[:8]
            year = entry.get("year", 2024)
            amount = entry.get("amount", 0.0)
            dtype = entry.get("type", "Other")
            qualified = "是" if entry.get("qualified", True) else "否"

            reversals = entry.get("reversals", [])
            reversed_total = sum(r.get("amount", 0.0) for r in reversals)
            remaining = max(0.0, amount - reversed_total)

            sts = get_dtl_status(entry, calc_year)
            status_label = sts["label"]

            recapture_amt = 0.0
            if sts["status"] == "recaptured":
                recapture_amt = remaining

            schedule = build_dtl_schedule(entry, calc_year)
            tl_parts = []
            for s in schedule:
                icon = {"created": "●", "reversal": "▼", "expiry": "◉",
                        "recapture": "✕", "recapture_warning": "⚠",
                        "cleared": "✓", "excluded": "○"}.get(s["milestone"], "")
                change = f" {s['change']:+,.0f}" if s["change"] != 0 else ""
                tl_parts.append(f"{icon}{s['year']}{change}")
            timeline = " → ".join(tl_parts)

            vals = [name, dtl_id, year, amount, dtype, qualified,
                    reversed_total, remaining, status_label,
                    sts.get("expiry_year", "—"), recapture_amt, timeline]
            for c, v in enumerate(vals, 1):
                ws.cell(row=rn, column=c, value=v)
            _style_body(ws, rn, len(headers))

            for c in (4, 7, 8, 11):
                cell = ws.cell(row=rn, column=c)
                cell.number_format = '#,##0.00'
                cell.font = NUMBER_FONT
                cell.alignment = RIGHT

            sts_color = {
                "active": "1F77B4", "near_expiry": "E6A817",
                "recaptured": "DC3545", "cleared": "28A745", "excluded": "6C757D",
            }.get(sts["status"], "333333")
            ws.cell(row=rn, column=9).font = Font(
                name="微软雅黑", size=10, bold=True, color=sts_color)

            jurisdiction_dtl_totals[name]["original"] += amount
            jurisdiction_dtl_totals[name]["reversed"] += reversed_total
            jurisdiction_dtl_totals[name]["remaining"] += remaining
            jurisdiction_dtl_totals[name]["recaptured"] += recapture_amt

            rn += 1

        if rn > jur_start_rn:
            jt = jurisdiction_dtl_totals[name]
            ws.merge_cells(f"A{rn}:C{rn}")
            ws.cell(row=rn, column=1, value=f"{name} 小计").font = BOLD_FONT
            for ci, v in [(4, jt["original"]), (7, jt["reversed"]),
                           (8, jt["remaining"]), (11, jt["recaptured"])]:
                cell = ws.cell(row=rn, column=ci, value=v)
                cell.number_format = '#,##0.00'
                cell.font = NUMBER_FONT
                cell.alignment = RIGHT
            for c in range(1, len(headers) + 1):
                ws.cell(row=rn, column=c).fill = SUBHEADER_FILL
                ws.cell(row=rn, column=c).border = Border(
                    bottom=Side(style="medium", color="1F4E79"))
            rn += 1

    _auto_width(ws, len(headers), max_width=50)
    _setup_print(ws)


# ═══════════════════════════════════════════════════════════
# Sheet 5: 合规追溯矩阵（NEW）
# ═══════════════════════════════════════════════════════════

def _write_compliance_matrix(ws, results: list[dict], rows: list[dict]) -> None:
    """OECD 条款合规追溯矩阵 —— 展示每个辖区每步计算对应的法条依据。"""
    ws.merge_cells("A1:F1")
    ws["A1"] = "OECD GloBE 规则合规追溯矩阵"
    ws["A1"].font = SECTION_FONT

    ws.merge_cells("A2:F2")
    ws["A2"] = "每条计算步骤均自动记录其适用的 OECD 条款——非 AI 匹配，而是代码执行时自记录。"
    ws["A2"].font = SMALL_FONT

    headers = ["辖区", "计算步骤", "OECD 条款", "详情", "状态"]
    hr = 4
    for c, h in enumerate(headers, 1):
        ws.cell(row=hr, column=c, value=h)
    _style_header(ws, hr, len(headers))

    rn = hr + 1
    STATUS_FILLS = {
        "applied": PatternFill(start_color="E8F5E9", end_color="E8F5E9", fill_type="solid"),
        "triggered": PatternFill(start_color="FFE8E8", end_color="FFE8E8", fill_type="solid"),
        "exempt": PatternFill(start_color="D4EDDA", end_color="D4EDDA", fill_type="solid"),
        "not_applicable": PatternFill(start_color="F5F5F5", end_color="F5F5F5", fill_type="solid"),
    }

    for r, row_data in zip(results, rows):
        name = row_data.get("name", "").strip()
        if not name:
            continue
        traces = r.get("_traces", [])
        if not traces:
            continue

        jur_start = rn
        for trace in traces:
            vals = [
                name,
                trace["step"],
                trace["article"],
                trace["detail"],
                TRACE_STATUS_ICONS.get(trace["status"], trace["status"]),
            ]
            for c, v in enumerate(vals, 1):
                ws.cell(row=rn, column=c, value=v)
            _style_body(ws, rn, len(headers))

            # 状态着色
            status_fill = STATUS_FILLS.get(trace["status"])
            if status_fill:
                ws.cell(row=rn, column=5).fill = status_fill
            if trace["status"] == "triggered":
                ws.cell(row=rn, column=5).font = Font(name="微软雅黑", size=10, bold=True, color="DC3545")

            rn += 1

        # 合并辖区列
        if rn > jur_start + 1:
            ws.merge_cells(f"A{jur_start}:A{rn - 1}")
            ws.cell(row=jur_start, column=1).alignment = Alignment(
                horizontal="center", vertical="center", wrap_text=True)
        elif rn == jur_start + 1:
            ws.cell(row=jur_start, column=1).alignment = CENTER

    _auto_width(ws, len(headers), min_width=14, max_width=80)
    ws.column_dimensions["D"].width = 65
    _setup_print(ws)


# ═══════════════════════════════════════════════════════════
# 主入口
# ═══════════════════════════════════════════════════════════

def export_gir(
    results: list[dict],
    rows: list[dict],
    allocation: dict,
    tax_flow: dict | None,
    sbie_year: int,
    group_name: str = "",
    etr_bar_fig=None,       # Plotly figure for ETR bar chart
    waterfall_fig=None,     # Plotly figure for Covered Taxes waterfall
    sankey_fig=None,        # Plotly figure for Sankey tax flow
    tax_analysis: dict | None = None,   # 云端税务分析（可选，附在报告末尾）
    result_review: dict | None = None,  # 本地结果复核结论（可选）
) -> io.BytesIO:
    """生成 GIR 标准格式 Excel 文件（5 Sheets + 可选云端分析 + 图表嵌入 + 合规追溯矩阵）。

    Sheet 顺序：
        1. 封面与摘要          — 基本信息 + 关键指标 + 辖区一览 + ETR柱状图
        2. ETR 计算与补税分配   — 计算明细 + 分配表 + 瀑布图
        3. Safe Harbour 与税源  — SH明细 + 税源流向 + Sankey全景图
        4. DTL 台账             — 递延税回转明细 + 辖区小计
        5. 合规追溯矩阵         — OECD 条款执行轨迹
        6. 云端分析结论         — LLM 分析与本地复核（可选，非计税依据）

    注意：Sheet 6 是分析性内容，**不参与任何计税**；所有金额仍由本地 calculator.py
    计算并落在 Sheet 1–5。
    """
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    group = group_name.strip() or (rows[0].get("name", "") if rows else "")
    total_topup = allocation.get("total_topup", 0.0)
    sh_count = sum(1 for r in results if r.get("safe_harbour"))
    high_count = sum(1 for r in results if r.get("risk") == "high")

    total_retained = 0.0
    total_exported = 0.0
    if tax_flow:
        total_retained = tax_flow.get("total_retained", 0.0)
        total_exported = tax_flow.get("total_exported", 0.0)

    # Sheet 1: 封面与摘要
    ws1 = wb.create_sheet("封面与摘要")
    _write_cover_summary(
        ws1, group, sbie_year, results, rows, allocation, tax_flow,
        total_topup, sh_count, high_count, total_retained, total_exported,
        etr_bar_fig,
    )

    # Sheet 2: ETR 计算与补税分配
    ws2 = wb.create_sheet("ETR计算与分配")
    _write_etr_and_allocation(ws2, results, rows, allocation, waterfall_fig)

    # Sheet 3: Safe Harbour 与税源流向
    ws3 = wb.create_sheet("SafeHarbour与税源")
    _write_sh_and_taxflow(ws3, results, rows, sbie_year, tax_flow, sankey_fig)

    # Sheet 4: DTL 台账
    ws4 = wb.create_sheet("DTL台账")
    _write_dtl(ws4, rows, sbie_year)

    # Sheet 5: 合规追溯矩阵
    ws5 = wb.create_sheet("合规追溯矩阵")
    _write_compliance_matrix(ws5, results, rows)

    # Sheet 6: 云端分析结论（可选，仅当有云端分析或本地复核结论时生成）
    if tax_analysis or result_review:
        ws6 = wb.create_sheet("云端分析结论")
        _write_analysis_sheet(ws6, tax_analysis, result_review)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _write_analysis_sheet(ws, tax_analysis: dict | None,
                          result_review: dict | None) -> None:
    """写入分析性内容：本地复核结论 + 云端税务分析。

    本表不参与计税，所有金额以 Sheet 1–5 的本地计算结果为准。
    """
    header_font = Font(bold=True, size=12)
    note_font = Font(italic=True, color="666666")

    row = 1
    ws.cell(row=row, column=1, value="分析结论（不参与计税）").font = Font(
        bold=True, size=14)
    row += 1
    ws.cell(row=row, column=1,
            value="本表内容为分析与解读，所有金额以「封面与摘要」等本地计算表为准。").font = note_font
    row += 2

    # 本地结果复核
    if result_review:
        ws.cell(row=row, column=1, value="一、本地结果复核（确定性）").font = header_font
        row += 1
        checks = result_review.get("checks") or []
        passed = sum(1 for c in checks if c.get("status") == "passed")
        ws.cell(row=row, column=1, value="审查结论")
        ws.cell(row=row, column=2, value=str(result_review.get("decision", "")))
        row += 1
        ws.cell(row=row, column=1, value="检查项")
        ws.cell(row=row, column=2, value=f"{passed}/{len(checks)} 项通过")
        row += 1
        ws.cell(row=row, column=1, value="错误 / 警告")
        ws.cell(row=row, column=2,
                value=f"{len(result_review.get('errors') or [])} / "
                      f"{len(result_review.get('warnings') or [])}")
        row += 1
        ws.cell(row=row, column=1, value="摘要")
        ws.cell(row=row, column=2, value=str(result_review.get("summary", "")))
        row += 2

    # 云端税务分析
    if tax_analysis:
        ws.cell(row=row, column=1, value="二、云端税务分析（LLM 生成）").font = header_font
        row += 1
        analysis = tax_analysis.get("analysis")
        if analysis:
            ws.cell(row=row, column=1, value="分析")
            ws.cell(row=row, column=2, value=str(analysis))
            row += 1
        for label, key in (("要点", "highlights"), ("行动建议", "actions")):
            values = tax_analysis.get(key)
            if not values:
                continue
            ws.cell(row=row, column=1, value=label)
            if isinstance(values, list):
                for item in values:
                    ws.cell(row=row, column=2, value=str(item))
                    row += 1
            else:
                ws.cell(row=row, column=2, value=str(values))
                row += 1
    else:
        ws.cell(row=row, column=1, value="本次运行未启用云端 AI，无云端分析内容。").font = note_font
        row += 1

    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 110


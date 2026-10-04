# -*- coding: utf-8 -*-
"""Phase 2 测试：云端建议比对本地映射、分析写入导出报告。

安全边界：云端建议只登记差异供人工确认，**不得覆盖本地确定性映射结果**。
"""
import io
import json

import openpyxl
import pytest

from Agent.agents import DataAgent
from Agent.agents.chart_agent import ChartAgent
from Agent.schemas import WorkflowState
from Agent.tools import export_gir_workbook


def _preview(*items):
    base = [
        {"globe_field": "profit", "label": "GloBE利润", "status": "matched",
         "status_label": "✅已匹配", "amount": 1000.0,
         "科目_detail": [{"科目": "利润总额", "金额": 1000.0, "matched": True}]},
        {"globe_field": "current_tax", "label": "当期所得税", "status": "matched",
         "status_label": "✅已匹配", "amount": 100.0,
         "科目_detail": [{"科目": "所得税费用", "金额": 100.0, "matched": True}]},
    ]
    return base + [dict(item) for item in items]


def _agent():
    return DataAgent(brain=None)


# ── 云端建议 vs 本地映射 ──

def test_no_comparison_without_suggestions():
    assert _agent()._compare_suggestions(None, _preview()) is None
    assert _agent()._compare_suggestions({"suggestions": []}, _preview()) is None
    assert _agent()._compare_suggestions("自由文本", _preview()) is None


def test_agreement_produces_no_review_items():
    suggestions = {"suggestions": [
        {"field": "profit", "source": "利润总额", "reason": "一致"}]}
    result = _agent()._compare_suggestions(suggestions, _preview())
    assert result["needs_review"] == 0
    assert result["agreed"] == 1
    assert result["changes_mapping"] is False
    assert "0 项需要人工确认" in result["summary"]


def test_conflicting_suggestion_is_flagged_not_applied():
    """本地已匹配、云端指向别的科目 → 登记为冲突，但不改本地结果。"""
    suggestions = {"suggestions": [
        {"field": "profit", "source": "净利润", "reason": "云端偏好净利润"}]}
    result = _agent()._compare_suggestions(suggestions, _preview())
    assert result["needs_review"] == 1
    conflict = result["conflicts"][0]
    assert conflict["field"] == "profit"
    assert conflict["local_sources"] == ["利润总额"]
    assert conflict["cloud_source"] == "净利润"
    assert result["changes_mapping"] is False


def test_locally_uncovered_field_becomes_pending_suggestion():
    suggestions = {"suggestions": [
        {"field": "payroll", "source": "管理费用", "reason": "本地没找到"}]}
    preview = _preview({"globe_field": "payroll", "label": "合格薪酬",
                        "status": "unmatched", "status_label": "⚠️未找到报表",
                        "amount": None, "科目_detail": []})
    result = _agent()._compare_suggestions(suggestions, preview)
    assert result["needs_review"] == 1
    item = result["uncovered"][0]
    assert item["field"] == "payroll"
    assert item["cloud_source"] == "管理费用"
    assert result["conflicts"] == []


def test_partial_field_is_treated_as_uncovered():
    suggestions = {"suggestions": [
        {"field": "tangible_assets", "source": "在建工程"}]}
    preview = _preview({"globe_field": "tangible_assets", "label": "有形资产",
                        "status": "partial", "status_label": "⚠️部分匹配",
                        "amount": 100.0, "科目_detail": []})
    result = _agent()._compare_suggestions(suggestions, preview)
    assert result["uncovered"][0]["field"] == "tangible_assets"


def test_fields_outside_rules_counted_as_unknown():
    suggestions = {"suggestions": [
        {"field": "made_up", "source": "无"}]}
    result = _agent()._compare_suggestions(suggestions, _preview())
    assert result["unknown_fields"] == 1
    assert result["needs_review"] == 0
    assert "规则外字段" in result["summary"]


def test_malformed_suggestions_are_skipped():
    suggestions = {"suggestions": [
        "不是对象", {"reason": "缺少 field"}, {"field": "  "},
        {"field": "profit", "source": "利润总额"}]}
    result = _agent()._compare_suggestions(suggestions, _preview())
    assert result["agreed"] == 1
    assert result["needs_review"] == 0


# ── 导出报告包含分析 ──

def _gir_bytes(tax_analysis=None, result_review=None):
    """用真实计算结果导出，避免手写 allocation 结构失真。"""
    from Agent.tools import calculate_rows

    rows = [{"name": "测试辖区", "profit": 1000.0, "current_tax": 100.0,
             "deferred_tax": 0.0, "revenue": 5000.0, "payroll": 0.0,
             "tangible_assets": 0.0, "parent_idx": None, "ownership": 1.0,
             "qdmtt_applies": False, "utpr_applies": True, "dtl_ledger": []}]
    calc = calculate_rows([dict(r) for r in rows], 2024)
    assert calc.ok, calc.error
    result = export_gir_workbook(calc.data["results"], rows, calc.data["allocation"],
                                 calc.data["tax_flow"], 2024,
                                 tax_analysis=tax_analysis,
                                 result_review=result_review)
    assert result.ok, result.error
    return result.data


def test_export_includes_analysis_sheet():
    analysis = {"analysis": "测试分析", "highlights": ["要点一"], "actions": ["建议一"]}
    review = {"decision": "pass", "summary": "复核通过", "checks": [
        {"check": "ETR", "status": "passed", "detail": ""}], "errors": [], "warnings": []}
    data = _gir_bytes(analysis, review)
    wb = openpyxl.load_workbook(io.BytesIO(data))
    assert "云端分析结论" in wb.sheetnames
    assert wb.sheetnames[-1] == "云端分析结论", "分析表应附在最后"

    text = "\n".join(
        str(cell) for row in wb["云端分析结论"].iter_rows(values_only=True)
        for cell in row if cell is not None)
    assert "不参与计税" in text
    assert "测试分析" in text
    assert "要点一" in text
    assert "建议一" in text
    assert "复核通过" in text


def test_export_states_when_cloud_disabled():
    data = _gir_bytes(tax_analysis=None, result_review={"decision": "pass",
                                                       "checks": [], "errors": [],
                                                       "warnings": [], "summary": "x"})
    wb = openpyxl.load_workbook(io.BytesIO(data))
    text = "\n".join(
        str(cell) for row in wb["云端分析结论"].iter_rows(values_only=True)
        for cell in row if cell is not None)
    assert "未启用云端" in text


def test_export_without_any_analysis_has_no_extra_sheet():
    data = _gir_bytes()
    wb = openpyxl.load_workbook(io.BytesIO(data))
    assert "云端分析结论" not in wb.sheetnames
    assert wb.sheetnames == ["封面与摘要", "ETR计算与分配", "SafeHarbour与税源",
                             "DTL台账", "合规追溯矩阵"]


def test_export_numbers_unaffected_by_analysis():
    """加入分析内容不得改变任何计算表内容。"""
    plain = _gir_bytes()
    with_analysis = _gir_bytes({"analysis": "x"}, None)
    wb_a = openpyxl.load_workbook(io.BytesIO(plain))
    wb_b = openpyxl.load_workbook(io.BytesIO(with_analysis))
    for sheet in ("封面与摘要", "ETR计算与分配", "SafeHarbour与税源", "合规追溯矩阵"):
        rows_a = list(wb_a[sheet].iter_rows(values_only=True))
        rows_b = list(wb_b[sheet].iter_rows(values_only=True))
        assert rows_a == rows_b, f"{sheet} 内容被分析表影响"

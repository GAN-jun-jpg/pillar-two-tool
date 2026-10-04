# -*- coding: utf-8 -*-
"""Agent tools 封装测试。"""
import io

import openpyxl

from Agent.tools import (
    build_charts,
    calculate_rows,
    export_gir_workbook,
    export_scenario_workbook,
    map_parsed_data,
    parse_financial_statement_bytes,
    validate_rows,
)


def _make_workbook_bytes() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "利润表"
    ws.append(["科目", "2025", "单位：万元"])
    ws.append(["营业总收入", 50000, ""])
    ws.append(["利润总额", 10000, ""])
    ws.append(["所得税费用", 1500, ""])
    ws.append(["递延所得税费用", 200, ""])

    ws2 = wb.create_sheet("资产负债表")
    ws2.append(["科目", "2025", "单位：万元"])
    ws2.append(["固定资产", 20000, ""])
    ws2.append(["资产总计", 50000, ""])

    ws3 = wb.create_sheet("现金流量表")
    ws3.append(["科目", "2025", "单位：万元"])
    ws3.append(["经营活动产生的现金流量净额", 8000, ""])
    ws3.append(["投资活动产生的现金流量净额", -2000, ""])
    ws3.append(["支付给职工以及为职工支付的现金", 1200, ""])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _row(**kw):
    base = {
        "name": "测试辖区", "profit": 1000.0, "current_tax": 100.0,
        "deferred_tax": 0.0, "revenue": 5000.0, "payroll": 0.0,
        "tangible_assets": 0.0, "parent_idx": None, "ownership": 1.0,
        "qdmtt_applies": False, "utpr_applies": True, "dtl_ledger": [],
    }
    base.update(kw)
    return base


def test_parse_tool():
    result = parse_financial_statement_bytes(_make_workbook_bytes(), "映射验证.xlsx")
    assert result.ok is True
    assert result.tool_name == "parse_financial_statement"
    assert result.data["detected_unit"] == "wan_yuan"
    assert result.data["sheets"]["cash_flow"] is not None


def test_map_tool():
    parsed = parse_financial_statement_bytes(_make_workbook_bytes(), "映射验证.xlsx")
    result = map_parsed_data(parsed.data, "测试辖区", unit=parsed.data["detected_unit"])
    assert result.ok is True
    rows = result.data["rows"]
    assert rows[0]["payroll"] == 1200.0
    assert rows[0]["current_tax"] == 1300.0
    assert rows[0]["tangible_assets"] == 20000.0


def test_validate_tool():
    result = validate_rows([_row()], calc_year=2024)
    assert result.ok is True
    assert result.data["has_errors"] is False


def test_calc_tool():
    result = calculate_rows([_row()], calc_year=2024)
    assert result.ok is True
    assert len(result.data["results"]) == 1
    assert result.data["summary"]["total_jurisdictions"] == 1
    assert result.data["summary"]["total_topup_tax"] == 50.0


def test_chart_tool():
    calc = calculate_rows([_row()], calc_year=2024)
    result = build_charts(calc.data["results"], [_row()], calc.data["tax_flow"])
    assert result.ok is True
    assert result.data["etr_bar"] is not None


def test_export_tools():
    rows = [_row()]
    scenario = export_scenario_workbook("测试方案", 2024, 0.10, 0.08, rows)
    assert scenario.ok is True
    assert scenario.data[:2] == b"PK"

    calc = calculate_rows(rows, calc_year=2024)
    gir = export_gir_workbook(
        calc.data["results"], rows, calc.data["allocation"],
        calc.data["tax_flow"], 2024, group_name="测试集团",
    )
    assert gir.ok is True
    assert gir.data[:2] == b"PK"


def test_tool_failure_returns_result():
    result = parse_financial_statement_bytes(b"not an excel file", "bad.xlsx")
    assert result.ok is False
    assert result.error
